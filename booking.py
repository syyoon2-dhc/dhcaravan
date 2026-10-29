# -*- coding: utf-8 -*-
"""동해카라반펜션 예약관리 모듈 (server.py 에서 불러서 사용)

데이터 위치 (모두 private/ 폴더 — 인터넷/GitHub 에 올라가지 않음)
  private/booking.db             예약 · 문자 발송 기록 (SQLite)
  private/booking_settings.json  입금계좌, 입금기한, 알리고 문자 설정, 문자 문구

요금 · 성수기 기간 · 공휴일은 data/prices.json 을 그대로 사용합니다.
"""
import io
import json
import os
import random
import re
import sqlite3
import string
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta

ROOT = os.path.dirname(os.path.abspath(__file__))
PRIVATE_DIR = os.environ.get('DHC_PRIVATE_DIR') or os.path.join(ROOT, 'private')
DB_FILE = os.path.join(PRIVATE_DIR, 'booking.db')
SETTINGS_FILE = os.path.join(PRIVATE_DIR, 'booking_settings.json')
PRICES_FILE = os.path.join(ROOT, 'data', 'prices.json')

ACTIVE = ('pending', 'confirmed', 'block')  # 객실을 차지하는 상태
STATUS_LABEL = {'pending': '입금대기', 'confirmed': '예약확정', 'cancelled': '취소', 'block': '예약막음'}
SOURCES = ('홈페이지', '전화', '현장', '땡큐캠핑', '기타')
PAY_METHODS = ('무통장', '카드', '현금', '땡큐캠핑', '기타')
CANCEL_BY = {'user': '사용자취소', 'admin': '관리자취소', 'auto': '입금기한초과'}
WEEKDAY = '월화수목금토일'

DEFAULT_SETTINGS = {
    'bookingOpen': True,           # 홈페이지 예약 받기
    'bank': '',                    # 입금 계좌 (예: 농협 000-0000-0000-00 홍길동) — 비어 있으면 예약 안 받음
    'bankShort': '',               # 문자용 짧은 계좌 (예: 기업 15912505001013 윤석영) — 비우면 위 계좌 사용
    'depositHours': 3,             # 예약 후 입금 기한 (시간)
    'maxDaysAhead': 180,           # 며칠 뒤까지 예약 받을지
    'openUntil': '',               # 이 날짜 숙박까지만 예약 오픈 (비우면 maxDaysAhead 기준)
    # 예약 오픈 방식: 'month' = 월단위 (땡큐캠핑과 같음) / 'date' = 마감일 지정 / 'days' = 오늘부터 며칠 / '' = 예전 방식
    # 월단위 예) 2개월 · 매월 1일 9시 → 10/1 09:00 부터 12/31 숙박까지 열림 (그 전에는 11/30 까지)
    'openMode': '', 'openMonths': 2, 'openDay': 1, 'openHour': 9,
    'maxNights': 6,                # 1회 최대 박수
    'maxRooms': 6,                 # 1회 최대 객실 수
    'phone': '010-4589-1415',      # 펜션 대표 번호 (문자 문구에 사용)
    'adminPhone': '010-4589-1415', # 새 예약 알림 받을 번호
    'adminNotify': True,
    'reminderHour': 10,            # 입실 전날 안내 문자 보내는 시각
    'units': {},                   # 객실별 동시 예약 가능 수 (기본 1) 예: {"deckzone": 3}
    'refundRules': [{'days': 6, 'rate': 100}, {'days': 5, 'rate': 90}, {'days': 4, 'rate': 80},  # 이용일 N일 전 이상 취소 → 환불 %
                    {'days': 3, 'rate': 70}, {'days': 2, 'rate': 50}, {'days': 1, 'rate': 30}, {'days': 0, 'rate': 0}],
    'nearDepositHours': 1,         # 입실 1일 전 · 당일 예약은 입금 기한 (시간)
    'nightDeposit': {'enabled': True, 'start': 21, 'end': 8},  # 밤 21시 ~ 아침 8시는 입금 기한 시간에서 뺌
    'depositReminderMinutes': 60,  # 입금 마감 N분 전 안내 문자
    'twoNight': {'enabled': True, 'weekdays': [4, 5], 'releaseDays': 7},  # 금(4)·토(5) 입실은 2박, 입실 7일 전부터 1박 허용
    'optionLimits': {},            # 옵션별 하루 최대 수량 {"1": 3} (옵션 번호: 수량)
    'specialPrices': [],           # 임시요금 [{rooms: ['gureum', ...] 또는 [] (전체), start, end, days: all|weekend|weekday, weekday, weekend, note}]
    'autoBlockCancels': 3,         # 연속 N회 취소한 고객은 자동으로 '예약불가' (0 = 사용 안 함)
    'refundNote': '환불은 취소 후 영업일 기준 3일 이내 입력하신 계좌로 입금됩니다.\n'
                  '천재지변 등 불가항력 시 별도 협의합니다.',
    'customerCancel': True,        # 손님이 예약조회에서 직접 취소 가능
    'smsPrice': 8.4,               # 알리고 단문 1건 요금 (원) — 비용 계산용
    'lmsPrice': 25,                # 알리고 장문 1건 요금 (원)
    'monthlyFixedCost': 0,         # 월 고정비 (도메인 · 서버 등, 원)
    'compareFeeRate': 7.7,         # 비교용: 땡큐캠핑 수수료율 (%, PG + 이용수수료)
    'customerCategories': ['일반', '단골', 'VIP', '예약불가'],  # 고객분류 ('예약불가' = 홈페이지 예약 차단)
    'notice': '입실 15:00 / 퇴실 11:00\n매너타임 22:00 ~ 07:00 (고성방가 · 소음 자제)\n'
              '카라반 · 객실 내부는 금연입니다.\n인원 추가는 영·유아를 포함한 인원 기준입니다.',
    'allowOrigins': ['https://syyoon2-dhc.github.io'],
    'aligo': {'enabled': False, 'key': '', 'userId': '', 'sender': '', 'testmode': False},
    'telegram': {'enabled': False, 'botToken': '', 'chatId': ''},  # 사장님 알림은 텔레그램으로 (무료)
    'auto': {'received': True, 'confirmed': True, 'cancelled': True, 'expired': True, 'reminder': True,
             'refund_done': True, 'deposit_reminder': True},
    'templates': {  # 단문(SMS 90바이트 = 한글 약 45자) 기준. 넘으면 장문(LMS)으로 발송
        'received': '[동해카라반]{금액}원 {계좌} {입금기한}까지 입금시 예약확정',
        'deposit_reminder': '[동해카라반]입금마감 1시간전입니다 {금액}원 {계좌}',
        'confirmed': '[동해카라반]{이름}님 {입실} {객실} 예약확정\n입실15시 퇴실11시 문의{펜션전화}',
        'expired': '[동해카라반]입금기한이 지나 {입실} 예약이 자동취소되었습니다',
        'cancelled': '[동해카라반]{입실} {객실} 예약이 취소되었습니다 문의{펜션전화}',
        'refund_done': '[동해카라반]{이름}님 환불금 {환불}원 입금완료되었습니다',
        'reminder': '[동해카라반]내일 입실일입니다. 15시부터 입실가능. 울주군 온산읍 강양길162',
        'admin': '[새예약]{이름} {연락처} {객실} {입실} {박}박 {금액}원',
        'admin_cancel': '[손님취소]{이름} {연락처} {입실} {객실} 환불{환불}원',
        'admin_expired': '[미입금 자동취소]{이름} {연락처} {객실} {입실} {금액}원',
    },
    # 사장님 텔레그램 알림 문구 (운영관리 → 문자발송설정에서 수정). 값이 비어 있는 항목이 든 줄은 자동으로 빠짐
    'tgTemplates': {
        'admin': '🔔 새 예약 (입금대기)\n{이름} {연락처}\n{입실} ~ {퇴실} ({박}박)\n{객실상세}\n합계 {금액}원\n'
                 '입금기한 {입금기한}\n요청사항: {요청사항}\n예약번호 {예약번호}',
        'admin_cancel': '↩️ 손님 직접 취소\n{이름} {연락처}\n{입실} ~ {퇴실} ({박}박)\n{객실상세}\n합계 {금액}원\n'
                        '환불할 금액 {환불}원\n환불계좌: {환불계좌}\n예약번호 {예약번호}',
        'admin_expired': '⏰ 미입금 자동취소\n{이름} {연락처}\n{입실} ~ {퇴실} ({박}박)\n{객실상세}\n합계 {금액}원\n예약번호 {예약번호}',
        'admin_paid': '💰 입금 확인 (예약확정)\n{이름} {연락처}\n{입실} ~ {퇴실} ({박}박)\n{객실상세}\n합계 {금액}원\n'
                      '입금액 {입금액}원 ({결제방법})\n잔금 {잔금}원\n예약번호 {예약번호}',
    },
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS reservations (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  code TEXT UNIQUE NOT NULL,
  group_code TEXT NOT NULL DEFAULT '',
  room_id TEXT NOT NULL,
  checkin TEXT NOT NULL,
  checkout TEXT NOT NULL,
  guests INTEGER NOT NULL DEFAULT 0,
  name TEXT NOT NULL DEFAULT '',
  phone TEXT NOT NULL DEFAULT '',
  request TEXT NOT NULL DEFAULT '',
  memo TEXT NOT NULL DEFAULT '',
  options TEXT NOT NULL DEFAULT '[]',
  room_amount INTEGER NOT NULL DEFAULT 0,
  extra_amount INTEGER NOT NULL DEFAULT 0,
  option_amount INTEGER NOT NULL DEFAULT 0,
  total INTEGER NOT NULL DEFAULT 0,
  paid INTEGER NOT NULL DEFAULT 0,
  refund INTEGER NOT NULL DEFAULT 0,
  status TEXT NOT NULL,
  source TEXT NOT NULL DEFAULT '홈페이지',
  deadline TEXT,
  cancel_reason TEXT NOT NULL DEFAULT '',
  reminded INTEGER NOT NULL DEFAULT 0,
  pay_method TEXT NOT NULL DEFAULT '무통장',
  paid_at TEXT,
  cancel_by TEXT NOT NULL DEFAULT '',
  cancelled_at TEXT,
  refund_status TEXT NOT NULL DEFAULT '',
  refund_bank TEXT NOT NULL DEFAULT '',
  refund_account TEXT NOT NULL DEFAULT '',
  refund_holder TEXT NOT NULL DEFAULT '',
  refund_done_at TEXT,
  deposit_reminded INTEGER NOT NULL DEFAULT 0,
  agreed_at TEXT,
  agree_ver TEXT NOT NULL DEFAULT '',
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_res_dates ON reservations(room_id, checkin, checkout);
CREATE INDEX IF NOT EXISTS idx_res_group ON reservations(group_code);
CREATE TABLE IF NOT EXISTS customers (
  phone TEXT PRIMARY KEY,
  name TEXT NOT NULL DEFAULT '',
  nickname TEXT NOT NULL DEFAULT '',
  category TEXT NOT NULL DEFAULT '',
  memo TEXT NOT NULL DEFAULT '',
  deleted INTEGER NOT NULL DEFAULT 0,
  created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS sms_log (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  created_at TEXT NOT NULL,
  reservation_id INTEGER,
  phone TEXT NOT NULL,
  kind TEXT NOT NULL,
  msg TEXT NOT NULL,
  ok INTEGER NOT NULL,
  result TEXT NOT NULL DEFAULT ''
);
CREATE TABLE IF NOT EXISTS visit_daily (
  day TEXT PRIMARY KEY,
  visitors INTEGER NOT NULL DEFAULT 0,
  home_views INTEGER NOT NULL DEFAULT 0,
  reserve_views INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS visit_source (
  day TEXT NOT NULL,
  source TEXT NOT NULL,
  visitors INTEGER NOT NULL DEFAULT 0,
  PRIMARY KEY (day, source)
);
CREATE TABLE IF NOT EXISTS visit_seen (
  day TEXT NOT NULL,
  vid TEXT NOT NULL,
  PRIMARY KEY (day, vid)
);
"""


class BookingError(Exception):
    """손님/관리자에게 그대로 보여줄 오류"""


_lock = threading.RLock()
_conn = None


def db():
    global _conn
    if _conn is None:
        os.makedirs(PRIVATE_DIR, exist_ok=True)
        _conn = sqlite3.connect(DB_FILE, check_same_thread=False)
        _conn.row_factory = sqlite3.Row
        _conn.executescript(SCHEMA)
        _migrate(_conn)
    return _conn


NEW_COLUMNS = {  # 예전 DB 파일에 없는 칸은 자동으로 추가
    'pay_method': "TEXT NOT NULL DEFAULT '무통장'", 'paid_at': 'TEXT', 'cancel_by': "TEXT NOT NULL DEFAULT ''",
    'cancelled_at': 'TEXT', 'refund_status': "TEXT NOT NULL DEFAULT ''", 'refund_bank': "TEXT NOT NULL DEFAULT ''",
    'refund_account': "TEXT NOT NULL DEFAULT ''", 'refund_holder': "TEXT NOT NULL DEFAULT ''", 'refund_done_at': 'TEXT',
    'deposit_reminded': 'INTEGER NOT NULL DEFAULT 0',
    'agreed_at': 'TEXT', 'agree_ver': "TEXT NOT NULL DEFAULT ''",
    'no_sms': 'INTEGER NOT NULL DEFAULT 0',  # 1 이면 이 예약은 손님 자동 문자를 보내지 않음 (관리자 등록 시 선택)
}

PRIVACY_VER = '2026-10-23'  # 예약 화면 동의 문구 · privacy.html 을 고치면 날짜도 바꾸기
KEEP_YEARS = 5  # 전자상거래법: 계약 · 청약철회 · 대금결제 기록 5년 보관 후 파기


def _migrate(conn):
    have = {r['name'] for r in conn.execute('PRAGMA table_info(reservations)')}
    for col, typ in NEW_COLUMNS.items():
        if col not in have:
            conn.execute(f'ALTER TABLE reservations ADD COLUMN {col} {typ}')
    # 예약에 있는 전화번호를 고객 목록에 등록 (처음 한 번)
    conn.execute("INSERT OR IGNORE INTO customers (phone, name, created_at) "
                 "SELECT phone, name, MIN(created_at) FROM reservations WHERE phone != '' GROUP BY phone")
    conn.commit()


def now_str():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


# ──────────────────────────────────────────────
#  설정
# ──────────────────────────────────────────────
def load_settings():
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    if os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, encoding='utf-8') as f:
            saved = json.load(f)
        for k, v in saved.items():
            if isinstance(v, dict) and isinstance(s.get(k), dict) and k != 'units':
                s[k].update(v)
            else:
                s[k] = v
    for kind, old in OLD_TEMPLATES.items():  # 예전 기본 문구를 그대로 쓰는 중이면 새 기본 문구로 (직접 고친 문구는 그대로)
        if s['templates'].get(kind) == old:
            s['templates'][kind] = DEFAULT_SETTINGS['templates'][kind]
    return s


OLD_TEMPLATES = {'confirmed': '[동해카라반]{이름}님 {입실} {객실} 예약확정. 입실15시 퇴실11시'}


def save_settings(data):
    s = load_settings()
    for k in DEFAULT_SETTINGS:
        if k in data:
            if isinstance(DEFAULT_SETTINGS[k], dict) and k != 'units':
                s[k].update(data[k] or {})
            else:
                s[k] = data[k]
    s['depositHours'] = max(1, int(s.get('depositHours') or 3))
    s['maxDaysAhead'] = max(1, int(s.get('maxDaysAhead') or 180))
    s['maxNights'] = max(1, int(s.get('maxNights') or 6))
    s['maxRooms'] = max(1, int(s.get('maxRooms') or 6))
    s['reminderHour'] = min(23, max(0, int(s.get('reminderHour') or 10)))
    s['units'] = {k: max(1, int(v)) for k, v in (s.get('units') or {}).items() if str(v).strip()}
    if s.get('openUntil'):
        parse_date(s['openUntil'], '예약 오픈 마감일')
    s['openMode'] = s.get('openMode') if s.get('openMode') in ('month', 'date', 'days') else ''
    s['openMonths'] = min(12, max(1, int(s.get('openMonths') or 2)))
    s['openDay'] = min(28, max(1, int(s.get('openDay') or 1)))
    s['openHour'] = min(23, max(0, int(s.get('openHour') or 0)))
    if s['openMode'] == 'date' and not s.get('openUntil'):
        raise BookingError('예약 오픈 마감일을 입력하세요')
    rules = {}
    for r in s.get('refundRules') or []:
        try:
            rules[max(0, int(r['days']))] = min(100, max(0, int(r['rate'])))
        except (KeyError, ValueError, TypeError):
            continue
    s['refundRules'] = [{'days': d, 'rate': rules[d]} for d in sorted(rules, reverse=True)]
    s['nearDepositHours'] = max(1, int(s.get('nearDepositHours') or 1))
    nd = s.get('nightDeposit') or {}
    hour = lambda v, d: min(23, max(0, int(d if v in (None, '') else v)))
    s['nightDeposit'] = {'enabled': bool(nd.get('enabled')), 'start': hour(nd.get('start'), 21), 'end': hour(nd.get('end'), 8)}
    s['depositReminderMinutes'] = max(0, int(s.get('depositReminderMinutes') or 0))
    s['autoBlockCancels'] = max(0, int(s.get('autoBlockCancels') or 0))
    tn = s.get('twoNight') or {}
    s['twoNight'] = {'enabled': bool(tn.get('enabled')), 'releaseDays': max(0, int(tn.get('releaseDays') or 0)),
                     'weekdays': sorted({int(w) for w in tn.get('weekdays') or [] if 0 <= int(w) <= 6})}
    s['optionLimits'] = {str(k): int(v) for k, v in (s.get('optionLimits') or {}).items() if str(v).strip() and int(v) > 0}
    specials = []
    for sp in s.get('specialPrices') or []:
        st, en = parse_date(sp.get('start'), '임시요금 시작일'), parse_date(sp.get('end'), '임시요금 종료일')
        if en < st:
            raise BookingError('임시요금 종료일이 시작일보다 빠릅니다')
        days = sp.get('days') if sp.get('days') in ('weekend', 'weekday') else 'all'
        rooms = sp.get('rooms') if isinstance(sp.get('rooms'), list) else ([sp['room']] if sp.get('room') else [])
        # 요금 방식: fixed=금액 직접 입력 / weekday=각 객실의 평일 요금 / minus=기본 요금에서 X원 할인 / percent=X% 할인
        mode = sp.get('mode') if sp.get('mode') in ('weekday', 'minus', 'percent') else 'fixed'
        item = {'rooms': sorted({str(r) for r in rooms if r}), 'start': st.isoformat(), 'end': en.isoformat(), 'days': days,
                'mode': mode, 'amount': int(sp.get('amount') or 0) if mode in ('minus', 'percent') else 0,
                'weekday': int(sp.get('weekday') or 0) if mode == 'fixed' and days != 'weekend' else 0,
                'weekend': int(sp.get('weekend') or 0) if mode == 'fixed' and days != 'weekday' else 0,
                'note': str(sp.get('note', ''))[:50]}
        if mode == 'fixed' and ((days != 'weekend' and item['weekday'] <= 0) or (days != 'weekday' and item['weekend'] <= 0)):
            raise BookingError('임시요금의 요금을 입력하세요')
        if mode in ('minus', 'percent') and item['amount'] <= 0:
            raise BookingError('할인 금액(또는 %)을 입력하세요')
        if mode == 'percent' and item['amount'] >= 100:
            raise BookingError('할인율은 100% 미만으로 입력하세요')
        specials.append(item)
    s['specialPrices'] = sorted(specials, key=lambda x: x['start'])
    os.makedirs(PRIVATE_DIR, exist_ok=True)
    with open(SETTINGS_FILE, 'w', encoding='utf-8') as f:
        json.dump(s, f, ensure_ascii=False, indent=2)
    return s


def refund_policy_text(settings):
    rules = sorted(settings.get('refundRules') or [], key=lambda r: -r['days'])
    lines = []
    for k, r in enumerate(rules):
        if r['days'] == 0:
            when = '이용 당일' if k == 0 else (f'이용 {rules[k - 1]["days"] - 1}일 전 ~ 당일' if rules[k - 1]['days'] > 1 else '이용 당일')
        elif k == 0:
            when = f'이용 {r["days"]}일 전까지'
        else:
            hi = rules[k - 1]['days'] - 1
            when = f'이용 {r["days"]}일 전' if hi == r['days'] else f'이용 {r["days"]}~{hi}일 전'
        lines.append(f'{when}: ' + ('환불 불가' if r['rate'] == 0 else f'{r["rate"]}% 환불'))
    return '\n'.join(lines + [settings.get('refundNote', '')]).strip()


def refund_rate(settings, checkin):
    days = (date.fromisoformat(checkin) - date.today()).days
    for r in sorted(settings.get('refundRules') or [], key=lambda r: -r['days']):
        if days >= r['days']:
            return r['rate']
    return 0


# ──────────────────────────────────────────────
#  요금 계산
# ──────────────────────────────────────────────
def load_prices():
    with open(PRICES_FILE, encoding='utf-8') as f:
        return json.load(f)


def parse_won(s):
    """'30,000원' -> 30000, '무료' -> 0"""
    digits = re.sub(r'[^\d]', '', str(s))
    return int(digits) if digits else 0


def extra_person_fee(room):
    """인원 추가 요금 (1인 1박). extraPerson 값이 없으면 비고 문구에서 읽음"""
    if room.get('extraPerson') not in (None, ''):
        return int(room['extraPerson'])
    m = re.search(r'([\d,]+)\s*원', room.get('extra', ''))
    return parse_won(m.group(1)) if m else 0


def paid_options(prices):
    """예약 때 고를 수 있는 옵션 목록 (인덱스 포함) — 유료 옵션 + '무료'로 적은 옵션(예: 월광보합, 선착순 수량 제한용)"""
    return [dict(o, i=i, amount=parse_won(o.get('price')))
            for i, o in enumerate(prices.get('options', []))
            if parse_won(o.get('price')) > 0 or '무료' in str(o.get('price', ''))]


def day_info(prices, d):
    """(성수기 여부, 주말 요금 여부)
    주말 = 토요일 · 공휴일 전날 · 성수기 금요일 (비수기 금요일은 평일 요금)"""
    ds = d.isoformat()
    peak = any(p.get('start', '') <= ds <= p.get('end', '') for p in prices.get('peakPeriods', []))
    before_holiday = (d + timedelta(days=1)).isoformat() in set(prices.get('holidays', []))
    wd = d.weekday()
    weekend = wd == 5 or before_holiday or (peak and wd == 4)
    return peak, weekend


def find_special(room, d, specials, weekend):
    """그날 그 객실에 적용되는 임시요금 (없으면 None). 기간이 겹치면 먼저 등록된 것 우선"""
    ds = d.isoformat()
    for sp in specials:
        days = sp.get('days', 'all')  # 주말만 · 평일만이면 나머지 요일은 기본 요금
        if (days == 'weekend' and not weekend) or (days == 'weekday' and weekend):
            continue
        rooms = sp.get('rooms') if 'rooms' in sp else ([sp['room']] if sp.get('room') else [])  # 예전 형식(room 하나)도 지원
        if (not rooms or room['id'] in rooms) and sp['start'] <= ds <= sp['end']:
            return sp
    return None


def night_price(prices, room, d, specials=()):
    peak, weekend = day_info(prices, d)
    base = room['price']['peak' if peak else 'normal'][1 if weekend else 0]
    sp = find_special(room, d, specials, weekend)  # 임시요금
    if not sp:
        return base, peak, weekend
    mode = sp.get('mode', 'fixed')
    if mode == 'weekday':  # 주말에도 그 객실의 평일 요금 (객실마다 다른 요금 그대로)
        return room['price']['peak' if peak else 'normal'][0], peak, weekend
    if mode == 'minus':
        return max(0, base - sp['amount']), peak, weekend
    if mode == 'percent':
        return base * (100 - sp['amount']) // 100 // 10 * 10, peak, weekend  # 10원 단위 버림
    return sp['weekend' if weekend else 'weekday'], peak, weekend


def get_room(prices, room_id):
    for r in prices['rooms']:
        if r['id'] == room_id:
            return r
    raise BookingError('객실을 찾을 수 없습니다')


def parse_date(s, label='날짜'):
    try:
        return date.fromisoformat(str(s))
    except ValueError:
        raise BookingError(f'{label} 형식이 올바르지 않습니다')


def quote(room_id, checkin, checkout, guests, options=None, prices=None, max_nights=60):
    """요금 계산. options = [{'i': 옵션번호, 'qty': 수량}]"""
    prices = prices or load_prices()
    room = get_room(prices, room_id)
    ci, co = parse_date(checkin, '입실일'), parse_date(checkout, '퇴실일')
    nights_n = (co - ci).days
    if nights_n < 1:
        raise BookingError('퇴실일은 입실일 다음 날 이후여야 합니다')
    if nights_n > max_nights:
        raise BookingError(f'최대 {max_nights}박까지 예약할 수 있습니다')
    try:
        guests = int(guests or 0)
    except ValueError:
        guests = 0
    if guests < 1:
        raise BookingError('인원을 선택해 주세요')
    if guests > room['max']:
        raise BookingError(f'{room["name"]}은(는) 최대 {room["max"]}명까지 이용할 수 있습니다')

    nights = []
    specials = load_settings().get('specialPrices') or []
    for k in range(nights_n):
        d = ci + timedelta(days=k)
        p, peak, weekend = night_price(prices, room, d, specials)
        nights.append({'date': d.isoformat(), 'price': p, 'peak': peak, 'weekend': weekend})
    room_amount = sum(n['price'] for n in nights)
    extra_n = max(0, guests - room['std'])
    extra_amount = extra_n * extra_person_fee(room) * nights_n

    by_index = {o['i']: o for o in paid_options(prices)}
    chosen = []
    for o in options or []:
        try:
            qty, opt = int(o.get('qty') or 0), by_index.get(int(o.get('i', -1)))
        except (ValueError, TypeError, AttributeError):
            continue
        if opt and qty > 0:
            chosen.append({'i': opt['i'], 'name': opt['name'], 'detail': opt.get('detail', ''),
                           'price': opt['amount'], 'qty': min(qty, 5)})
    option_amount = sum(o['price'] * o['qty'] for o in chosen)

    return {
        'room_id': room_id, 'room_name': room['name'], 'checkin': ci.isoformat(), 'checkout': co.isoformat(),
        'nights': nights, 'guests': guests, 'extra_guests': extra_n,
        'room_amount': room_amount, 'extra_amount': extra_amount,
        'options': chosen, 'option_amount': option_amount,
        'total': room_amount + extra_amount + option_amount,
    }


# ──────────────────────────────────────────────
#  예약 가능 여부
# ──────────────────────────────────────────────
def units_of(settings, room_id):
    return int((settings.get('units') or {}).get(room_id, 1) or 1)


def booked_counts(room_id, start, end, exclude_id=None):
    """start~end(미포함) 각 날짜별로 객실을 차지한 예약 수"""
    rows = db().execute(
        f"SELECT id, checkin, checkout FROM reservations WHERE room_id=? AND status IN ({','.join('?' * len(ACTIVE))}) "
        "AND checkin < ? AND checkout > ?",
        (room_id, *ACTIVE, end.isoformat(), start.isoformat())).fetchall()
    counts = {}
    for r in rows:
        if exclude_id and r['id'] == exclude_id:
            continue
        d = max(date.fromisoformat(r['checkin']), start)
        last = min(date.fromisoformat(r['checkout']), end)
        while d < last:
            counts[d] = counts.get(d, 0) + 1
            d += timedelta(days=1)
    return counts


def is_available(settings, room_id, checkin, checkout, exclude_id=None):
    ci, co = date.fromisoformat(checkin), date.fromisoformat(checkout)
    counts = booked_counts(room_id, ci, co, exclude_id)
    units = units_of(settings, room_id)
    return all(counts.get(ci + timedelta(days=k), 0) < units for k in range((co - ci).days))


def _month_end(y, m):
    """y년 m월(13 이상이면 다음 해로 넘김)의 마지막 날"""
    y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
    return (date(y, m, 28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def month_open_info(settings, now=None):
    """월단위 오픈: (지금 열려 있는 마지막 숙박일, 다음 오픈 시각, 다음 오픈 때 열리는 마지막 숙박일)"""
    now = now or datetime.now()
    months, day, hour = settings.get('openMonths') or 2, settings.get('openDay') or 1, settings.get('openHour') or 0
    this_open = datetime(now.year, now.month, day, hour)
    if now >= this_open:  # 이번 달 오픈 시각이 지남 → 이번 달 + N개월 말일까지
        until = _month_end(now.year, now.month + months)
        nxt_y, nxt_m = now.year + now.month // 12, now.month % 12 + 1
        nxt = datetime(nxt_y, nxt_m, day, hour)
    else:  # 아직 이번 달 오픈 전 → 이번 달 + (N-1)개월 말일까지
        until = _month_end(now.year, now.month + months - 1)
        nxt = this_open
    return until, nxt, _month_end(until.year, until.month + 1)


def next_open_text(settings):
    """월단위 오픈일 때 다음 오픈 안내 (예: '11/1(일) 9시에 1월 예약이 열립니다')"""
    if settings.get('openMode') != 'month':
        return ''
    _, nxt, nxt_until = month_open_info(settings)
    return f'{fmt_day(nxt.date().isoformat())} {nxt.hour}시에 {nxt_until.month}월 예약이 열립니다'


def open_until(settings):
    """예약을 받는 마지막 숙박일"""
    mode = settings.get('openMode') or ''
    if mode == 'month':
        return month_open_info(settings)[0]
    if mode in ('date', '') and settings.get('openUntil'):
        try:
            return date.fromisoformat(settings['openUntil'])
        except ValueError:
            pass
    return date.today() + timedelta(days=settings['maxDaysAhead'])


# ──────────────────────────────────────────────
#  예약 저장 / 변경
# ──────────────────────────────────────────────
def new_code():
    while True:
        code = 'DH' + datetime.now().strftime('%y%m%d') + '-' + ''.join(random.choices(string.digits, k=4))
        if not db().execute('SELECT 1 FROM reservations WHERE code LIKE ? OR group_code=?',
                            (code + '%', code)).fetchone():
            return code


def norm_phone(p):
    return re.sub(r'\D', '', str(p or ''))


def fmt_phone(p):
    p = norm_phone(p)
    if len(p) == 11:
        return f'{p[:3]}-{p[3:7]}-{p[7:]}'
    if len(p) == 10:
        return f'{p[:3]}-{p[3:6]}-{p[6:]}'
    return p


def valid_mobile(p):
    return bool(re.fullmatch(r'01\d{8,9}', norm_phone(p)))


def row_dict(r):
    d = dict(r)
    d['options'] = json.loads(d.get('options') or '[]')
    d['status_label'] = STATUS_LABEL.get(d['status'], d['status'])
    d['nights'] = (date.fromisoformat(d['checkout']) - date.fromisoformat(d['checkin'])).days
    d['phone_fmt'] = fmt_phone(d['phone'])
    d['cancel_by_label'] = CANCEL_BY.get(d.get('cancel_by') or '', '')
    d['refund_label'] = {'due': '환불예정', 'done': '환불완료'}.get(d.get('refund_status') or '', '')
    d['refund_info'] = ' '.join(x for x in (d.get('refund_bank'), d.get('refund_account'), d.get('refund_holder')) if x)
    d['discount'] = d['room_amount'] + d['extra_amount'] + d['option_amount'] - d['total']
    return d


def get_reservation(rid):
    r = db().execute('SELECT * FROM reservations WHERE id=?', (rid,)).fetchone()
    if not r:
        raise BookingError('예약을 찾을 수 없습니다')
    return row_dict(r)


def group_rows(group_code):
    return [row_dict(r) for r in db().execute(
        'SELECT * FROM reservations WHERE group_code=? ORDER BY id', (group_code,)).fetchall()]


def _insert(q, status, source, deadline, name, phone, request, memo, total, paid, code, group_code, pay_method='무통장'):
    ts = now_str()
    cur = db().execute(
        'INSERT INTO reservations (code, group_code, room_id, checkin, checkout, guests, name, phone, request, memo, '
        'options, room_amount, extra_amount, option_amount, total, paid, status, source, deadline, pay_method, '
        'paid_at, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
        (code, group_code, q['room_id'], q['checkin'], q['checkout'], q['guests'], name, phone, request, memo,
         json.dumps(q['options'], ensure_ascii=False), q['room_amount'], q['extra_amount'], q['option_amount'],
         total, paid, status, source, deadline, pay_method, ts if paid else None, ts, ts))
    if phone:
        db().execute('INSERT OR IGNORE INTO customers (phone, name, created_at) VALUES (?,?,?)', (phone, name, ts))
    return cur.lastrowid


def deposit_deadline(settings, checkin, hours=None):
    """입금 기한: 기본 N시간 / 입실 1일 전 · 당일 예약은 1시간 / 입실일 자정을 넘으면 입실 전날 23:59
    밤 시간(기본 21시 ~ 아침 8시)은 시간에서 뺌 (사장님이 밤에는 입금 확인을 못 하므로)
    예: 3시간 기한 — 20:50 예약 → 21시까지 10분 + 다음날 8시부터 2시간 50분 = 다음날 10:50, 밤 예약 → 다음날 11:00"""
    now = datetime.now()
    ci = date.fromisoformat(checkin)
    if hours is None:
        hours = settings['nearDepositHours'] if (ci - now.date()).days <= 1 else settings['depositHours']
    deadline = add_day_hours(settings, now, int(hours))
    cap = datetime.combine(ci, datetime.min.time()) - timedelta(minutes=1)
    if deadline != now + timedelta(hours=int(hours)):  # 밤을 건너뛴 기한은 입실일 낮 12시까지 허용 (입실 전날 밤 예약도 아침에 입금)
        cap = datetime.combine(ci, datetime.min.time()) + timedelta(hours=12)
    if cap > now + timedelta(minutes=30):
        deadline = min(deadline, cap)
    return deadline.strftime('%Y-%m-%d %H:%M:%S')


def add_day_hours(settings, now, hours):
    """now 부터 hours 시간 뒤 — 밤 시간(nightDeposit start ~ 다음날 end)은 세지 않음"""
    nd = settings.get('nightDeposit') or {}
    a, b = int(nd.get('start', 21)), int(nd.get('end', 8))
    if not nd.get('enabled') or a <= b:
        return now + timedelta(hours=hours)
    t, left = now, timedelta(hours=hours)
    while True:
        morning = t.replace(hour=b, minute=0, second=0, microsecond=0)
        night = t.replace(hour=a, minute=0, second=0, microsecond=0)
        if t < morning:          # 새벽 → 그날 아침부터
            t = morning
        elif t >= night:         # 밤 → 다음날 아침부터
            t = morning + timedelta(days=1)
            continue
        if t + left <= night:
            return t + left
        left -= night - t
        t = morning + timedelta(days=1)


def check_two_night(settings, checkin, checkout):
    """2박 우선: 금 · 토 입실은 2박 이상 (입실 N일 전부터는 1박 가능)"""
    tn = settings.get('twoNight') or {}
    ci, co = date.fromisoformat(checkin), date.fromisoformat(checkout)
    if not tn.get('enabled') or (co - ci).days != 1 or ci.weekday() not in tn.get('weekdays', []):
        return
    if (ci - date.today()).days < tn.get('releaseDays', 0):
        return
    days = '·'.join('월화수목금토일'[w] for w in tn['weekdays'])
    raise BookingError(f'{days}요일 입실은 2박 이상 예약할 수 있습니다. '
                       f'(입실 {tn["releaseDays"]}일 전부터는 1박도 가능)')


def option_left(settings, checkin, checkout, exclude_id=None):
    """수량 제한이 있는 옵션의 남은 수량 {"0": 2} — 묵는 날 중 가장 적게 남은 날 기준 (exclude_id: 수정 중인 예약은 빼고 계산)"""
    limits = settings.get('optionLimits') or {}
    try:
        ci, co = date.fromisoformat(str(checkin)), date.fromisoformat(str(checkout))
    except ValueError:
        return {}
    if not limits or co <= ci:
        return {}
    rows = [row_dict(r) for r in db().execute(
        "SELECT * FROM reservations WHERE status IN ('pending','confirmed') AND checkin < ? AND checkout > ? "
        "AND options != '[]' AND id != ?", (co.isoformat(), ci.isoformat(), int(exclude_id or 0))).fetchall()]
    left = {}
    d = ci
    while d < co:
        ds = d.isoformat()
        for i, limit in limits.items():
            used = sum(o['qty'] for r in rows if r['checkin'] <= ds < r['checkout'] for o in r['options'] if str(o['i']) == i)
            left[i] = min(left.get(i, limit), max(0, limit - used))
        d += timedelta(days=1)
    return left


def check_option_limits(settings, quotes):
    """옵션 하루 수량 제한 (예: 바비큐 장비 하루 3개)"""
    limits = settings.get('optionLimits') or {}
    if not limits:
        return
    ci, co = quotes[0]['checkin'], quotes[0]['checkout']
    want = {}
    for q in quotes:
        for o in q['options']:
            want[str(o['i'])] = want.get(str(o['i']), 0) + o['qty']
    rows = [row_dict(r) for r in db().execute(
        "SELECT * FROM reservations WHERE status IN ('pending','confirmed') AND checkin < ? AND checkout > ? "
        "AND options != '[]'", (co, ci)).fetchall()]
    d = date.fromisoformat(ci)
    while d < date.fromisoformat(co):
        ds = d.isoformat()
        for i, qty in want.items():
            if i not in limits:
                continue
            used = sum(o['qty'] for r in rows if r['checkin'] <= ds < r['checkout'] for o in r['options']
                       if str(o['i']) == i)
            if used + qty > limits[i]:
                name = next(o['name'] + ' ' + o.get('detail', '') for q in quotes for o in q['options'] if str(o['i']) == i)
                left = max(0, limits[i] - used)
                raise BookingError(f'{fmt_day(ds)} {name}은(는) {left}개만 남았습니다')
        d += timedelta(days=1)


def check_open_range(settings, checkin, checkout):
    ci, co = parse_date(checkin, '입실일'), parse_date(checkout, '퇴실일')
    if ci < date.today():
        raise BookingError('지난 날짜는 예약할 수 없습니다')
    if co - timedelta(days=1) > open_until(settings):
        raise BookingError(f'현재 {fmt_day(open_until(settings).isoformat())} 숙박까지만 예약할 수 있습니다')


def create_web_booking(data):
    """홈페이지 예약 (객실 여러 개 한 번에). 모두 '입금대기'로 들어오고 하나의 예약번호로 묶임
    data = {checkin, checkout, name, phone, request, items: [{room_id, guests, options}]}"""
    settings, prices = load_settings(), load_prices()
    if not settings.get('bookingOpen', True) or not settings.get('bank'):
        raise BookingError('현재 온라인 예약을 받지 않고 있습니다. 전화로 문의해 주세요')
    name = str(data.get('name', '')).strip()[:30]
    phone = norm_phone(data.get('phone'))
    if not name:
        raise BookingError('예약자 이름을 입력해 주세요')
    if not valid_mobile(phone):
        raise BookingError('휴대폰 번호를 정확히 입력해 주세요')
    if is_blocked_customer(phone):
        raise BookingError(f'온라인 예약이 어려운 번호입니다. {settings["phone"]}로 문의해 주세요')
    items = data.get('items') or []
    if not isinstance(items, list) or not items:
        raise BookingError('객실을 선택해 주세요')
    if len(items) > settings['maxRooms']:
        raise BookingError(f'한 번에 최대 {settings["maxRooms"]}개 객실까지 예약할 수 있습니다')
    if len({it.get('room_id') for it in items}) != len(items):
        raise BookingError('같은 객실이 중복 선택되었습니다')
    check_open_range(settings, data.get('checkin'), data.get('checkout'))
    quotes = [quote(it.get('room_id'), data.get('checkin'), data.get('checkout'), it.get('guests'),
                    it.get('options'), prices, settings['maxNights']) for it in items]
    request = str(data.get('request', '')).strip()[:500]
    agree = data.get('agree') or {}
    if not (isinstance(agree, dict) and all(agree.get(k) is True for k in ('notice', 'refund', 'privacy'))):
        raise BookingError('필수 항목에 모두 동의해 주세요')
    check_two_night(settings, quotes[0]['checkin'], quotes[0]['checkout'])

    with _lock:
        expire_overdue()
        for q in quotes:
            if not is_available(settings, q['room_id'], q['checkin'], q['checkout']):
                raise BookingError(f'{q["room_name"]}은(는) 방금 다른 분이 예약했습니다. 다시 선택해 주세요')
        check_option_limits(settings, quotes)
        base = new_code()
        deadline = deposit_deadline(settings, quotes[0]['checkin'])
        for k, q in enumerate(quotes):
            code = base if len(quotes) == 1 else f'{base}-{k + 1}'
            _insert(q, 'pending', '홈페이지', deadline, name, phone, request, '', q['total'], 0, code, base)
        db().execute('UPDATE reservations SET agreed_at=?, agree_ver=? WHERE group_code=?',
                     (now_str(), PRIVACY_VER, base))  # 개인정보 수집 · 이용 동의 기록
        db().commit()
        rows = group_rows(base)

    notify('received', rows)
    if settings.get('adminNotify'):
        notify('admin', rows)
    return {'code': base, 'total': sum(r['total'] for r in rows), 'bank': settings['bank'],
            'deadline': fmt_deadline(deadline), 'phone': settings['phone']}


def create_reservation(data):
    """관리자가 직접 넣는 예약 (전화 · 현장 · 땡큐캠핑 이관 · 예약막음)
    data['extra_rooms'] = [{room_id, guests}] 이면 같은 날짜 · 같은 손님으로 객실 여러 개를 한 예약번호로 묶어 등록
    (결제금액 · 입금액 · 옵션은 첫 객실 기준, 추가 객실은 자동 계산 요금)"""
    settings, prices = load_settings(), load_prices()
    status = data.get('status', 'confirmed')
    if status not in ('pending', 'confirmed', 'block'):
        raise BookingError('잘못된 예약 상태')
    name = str(data.get('name', '')).strip()[:30]
    phone = norm_phone(data.get('phone'))
    source = data.get('source') if data.get('source') in SOURCES else '기타'
    pay_method = data.get('pay_method') if data.get('pay_method') in PAY_METHODS else '무통장'

    extra_rooms = [x for x in (data.get('extra_rooms') or []) if isinstance(x, dict) and x.get('room_id')]
    room_ids = [data.get('room_id')] + [x['room_id'] for x in extra_rooms]
    if len(set(room_ids)) != len(room_ids):
        raise BookingError('같은 객실이 중복 선택되었습니다')

    def block_quote(room_id):
        bq = quote(room_id, data.get('checkin'), data.get('checkout'), 1, [], prices)
        bq.update(room_amount=0, extra_amount=0, option_amount=0, total=0, options=[], guests=0)
        return bq

    if status == 'block':
        q = block_quote(data.get('room_id'))
        extra_qs = [block_quote(x['room_id']) for x in extra_rooms]
        name, phone, source = name or '예약막음', '', '기타'
    else:
        q = quote(data.get('room_id'), data.get('checkin'), data.get('checkout'), data.get('guests'),
                  data.get('options'), prices)
        extra_qs = [quote(x['room_id'], data.get('checkin'), data.get('checkout'),
                          x.get('guests') or get_room(prices, x['room_id'])['std'], [], prices) for x in extra_rooms]
        if not name:
            raise BookingError('예약자 이름을 입력해 주세요')
        if phone and not valid_mobile(phone):
            raise BookingError('휴대폰 번호를 정확히 입력해 주세요')

    total = q['total'] if data.get('total') in (None, '') else int(data['total'])  # 할인 등 직접 조정 가능
    if data.get('paid') not in (None, ''):
        paid = int(data['paid'])
    else:
        paid = total if status == 'confirmed' else 0

    plan = [(q, total, paid)] + [(eq, eq['total'], eq['total'] if status == 'confirmed' else 0) for eq in extra_qs]
    with _lock:
        expire_overdue()
        if not data.get('force'):
            for pq, _, _ in plan:
                if not is_available(settings, pq['room_id'], pq['checkin'], pq['checkout']):
                    raise BookingError(f'{pq["room_name"]}: 해당 날짜에 이미 예약이 있습니다')
        deadline = deposit_deadline(settings, q['checkin'], data.get('depositHours')) if status == 'pending' else None
        base = new_code()
        request, memo = str(data.get('request', '')).strip()[:500], str(data.get('memo', '')).strip()[:1000]
        ids = []
        for k, (pq, ptotal, ppaid) in enumerate(plan):
            code = base if len(plan) == 1 else f'{base}-{k + 1}'
            rid = _insert(pq, status, source, deadline, name, phone, request, memo, ptotal, ppaid, code, base, pay_method)
            extra = {'no_sms': 1 if data.get('no_sms') else 0}
            if ppaid and data.get('paid_date'):  # 통장에 실제로 입금된 날 (매출 · 입금 달력 기준)
                extra['paid_at'] = parse_date(data['paid_date'], '입금일').isoformat() + now_str()[10:]
            db().execute(f'UPDATE reservations SET {", ".join(k + "=?" for k in extra)} WHERE id=?', (*extra.values(), rid))
            ids.append(rid)
        db().commit()
        rows = [get_reservation(i) for i in ids]

    if data.get('sms') and phone and status in ('pending', 'confirmed'):
        notify('received' if status == 'pending' else 'confirmed', rows, force=True)
    return rows[0]


EDITABLE = ('room_id', 'checkin', 'checkout', 'guests', 'name', 'phone', 'request', 'memo', 'source',
            'room_amount', 'extra_amount', 'option_amount', 'total', 'paid', 'refund', 'deadline', 'options',
            'pay_method', 'refund_bank', 'refund_account', 'refund_holder', 'refund_status', 'no_sms')


def update_reservation(rid, data):
    settings = load_settings()
    with _lock:
        r = get_reservation(rid)
        new = {k: data[k] for k in EDITABLE if k in data}
        if 'phone' in new:
            new['phone'] = norm_phone(new['phone'])
            if new['phone'] and not valid_mobile(new['phone']):
                raise BookingError('휴대폰 번호를 정확히 입력해 주세요')
        if 'source' in new and new['source'] not in SOURCES:
            new['source'] = '기타'
        if 'pay_method' in new and new['pay_method'] not in PAY_METHODS:
            new['pay_method'] = '기타'
        if 'refund_status' in new and new['refund_status'] not in ('', 'due', 'done'):
            new['refund_status'] = ''
        if 'no_sms' in new:
            new['no_sms'] = 1 if new['no_sms'] else 0
        for k in ('guests', 'room_amount', 'extra_amount', 'option_amount', 'total', 'paid', 'refund'):
            if k in new:
                new[k] = int(new[k] or 0)
        if 'options' in new:
            new['options'] = json.dumps(new['options'], ensure_ascii=False)
        if 'deadline' in new and not new['deadline']:
            new['deadline'] = None
        room_id = new.get('room_id', r['room_id'])
        ci = parse_date(new.get('checkin', r['checkin']), '입실일').isoformat()
        co = parse_date(new.get('checkout', r['checkout']), '퇴실일').isoformat()
        if co <= ci:
            raise BookingError('퇴실일은 입실일 다음 날 이후여야 합니다')
        get_room(load_prices(), room_id)
        moved = (room_id, ci, co) != (r['room_id'], r['checkin'], r['checkout'])
        if moved and r['status'] in ACTIVE and not data.get('force'):
            if not is_available(settings, room_id, ci, co, exclude_id=rid):
                raise BookingError('변경하려는 날짜에 이미 다른 예약이 있습니다')
        if moved:
            new['reminded'] = 0
        if not new:
            return r
        new['updated_at'] = now_str()
        db().execute(f'UPDATE reservations SET {", ".join(k + "=?" for k in new)} WHERE id=?', (*new.values(), rid))
        db().commit()
        return get_reservation(rid)


def set_status(rid, status, data=None):
    """상태 변경 (입금확인 → 확정, 취소, 되돌리기).
    같은 예약번호로 묶인 객실(여러 객실 예약)은 data['group'] 이 참이면 함께 변경"""
    data = data or {}
    settings = load_settings()
    if status not in ('pending', 'confirmed', 'cancelled'):
        raise BookingError('잘못된 상태')
    with _lock:
        r = get_reservation(rid)
        if r['status'] == 'block':
            raise BookingError('예약막음은 삭제로 해제하세요')
        targets = [r]
        if data.get('group', True) and r['group_code']:
            targets = [x for x in group_rows(r['group_code']) if x['status'] == r['status']]
        single = len(targets) == 1
        for t in targets:
            fields = {'status': status, 'updated_at': now_str()}
            if status == 'confirmed':
                fields['paid'] = int(data['paid']) if single and data.get('paid') not in (None, '') else t['total']
                fields['deadline'] = None
                fields['paid_at'] = now_str()
                if data.get('pay_method') in PAY_METHODS:
                    fields['pay_method'] = data['pay_method']
            elif status == 'cancelled':
                refund = int(data.get('refund') or 0) if single else (t['paid'] if data.get('refundAll') else 0)
                fields.update(cancel_reason=str(data.get('reason', '')).strip()[:200] or '관리자 취소',
                              cancel_by='admin', cancelled_at=now_str(), refund=refund,
                              refund_status='due' if refund > 0 else '', deadline=None,
                              refund_bank=str(data.get('refund_bank', '')).strip()[:30],
                              refund_account=str(data.get('refund_account', '')).strip()[:40],
                              refund_holder=str(data.get('refund_holder', '')).strip()[:30])
            elif status == 'pending':
                fields['deadline'] = deposit_deadline(settings, t['checkin'])
                fields['deposit_reminded'] = 0
            if t['status'] == 'cancelled':
                if not is_available(settings, t['room_id'], t['checkin'], t['checkout'], exclude_id=t['id']):
                    raise BookingError('그 사이에 같은 날짜에 다른 예약이 들어와 복구할 수 없습니다')
                fields.update(cancel_reason='', refund=0, cancel_by='', cancelled_at=None, refund_status='',
                              refund_done_at=None)
            db().execute(f'UPDATE reservations SET {", ".join(k + "=?" for k in fields)} WHERE id=?',
                         (*fields.values(), t['id']))
        db().commit()
        rows = [get_reservation(t['id']) for t in targets]
    if data.get('sms') and status in ('confirmed', 'cancelled') and rows[0]['phone']:
        notify(status, rows, force=True)
    if status == 'confirmed' and r['status'] == 'pending':  # 입금대기 → 확정: 사장님 텔레그램에 입금확인 알림
        notify('admin_paid', rows)
    return rows


def delete_reservation(rid):
    with _lock:
        db().execute('DELETE FROM reservations WHERE id=?', (rid,))
        db().commit()


def list_reservations(start=None, end=None, status=None, q=None, basis='stay', room=None, refund_status=None,
                      cancel_by=None):
    """basis: stay=숙박 기간이 겹치는 예약 / checkin=입실일 / created=신청일 / cancelled=취소일 기준"""
    sql, args = 'SELECT * FROM reservations WHERE 1=1', []
    if start and end:
        if basis == 'stay':
            sql += ' AND checkin < ? AND checkout > ?'
            args += [end, start]
        else:
            col = {'checkin': 'checkin', 'created': 'created_at', 'cancelled': 'cancelled_at'}.get(basis, 'checkin')
            sql += f' AND {col} >= ? AND {col} < ?'
            args += [start, end]
    if room:
        sql += ' AND room_id=?'
        args.append(room)
    if refund_status:
        sql += f' AND refund_status IN ({",".join("?" * len(refund_status))})'
        args += list(refund_status)
    if cancel_by:
        sql += ' AND cancel_by=?'
        args.append(cancel_by)
    if status:
        sql += f' AND status IN ({",".join("?" * len(status))})'
        args += list(status)
    if q:
        like = f'%{q}%'
        sql += ' AND (name LIKE ? OR code LIKE ? OR memo LIKE ?'
        args += [like, like, like]
        if norm_phone(q):
            sql += ' OR phone LIKE ?'
            args.append(f'%{norm_phone(q)}%')
        sql += ')'
    order = {'created': 'created_at DESC', 'cancelled': 'cancelled_at DESC'}.get(basis, 'checkin, room_id')
    sql += f' ORDER BY {order} LIMIT 3000'
    with _lock:
        return [row_dict(r) for r in db().execute(sql, args).fetchall()]


def lookup(code, phone):
    """예약번호 또는 예약자 이름 + 휴대폰으로 조회. 이름으로 찾으면 다가오는 예약(없으면 최근 예약)"""
    key, phone = str(code or '').strip(), norm_phone(phone)
    rows = []
    with _lock:
        if key and phone:
            found = [row_dict(r) for r in db().execute(
                "SELECT * FROM reservations WHERE (group_code=? OR code=? OR name=?) AND phone=? AND status != 'block' "
                "ORDER BY id", (key.upper(), key.upper(), key, phone)).fetchall()]
            groups = {}
            for r in found:
                groups.setdefault(r['group_code'] or r['code'], []).append(r)
            if groups:
                today = date.today().isoformat()
                upcoming = sorted((g for g in groups.values() if g[0]['checkout'] > today and
                                   any(x['status'] != 'cancelled' for x in g)), key=lambda g: g[0]['checkin'])
                rows = upcoming[0] if upcoming else max(groups.values(), key=lambda g: g[0]['created_at'])
    if not rows:
        raise BookingError('예약 내역을 찾을 수 없습니다. 예약번호(또는 예약자 이름)와 휴대폰 번호를 확인해 주세요')
    names = {r['id']: r['name'] for r in load_prices()['rooms']}
    settings = load_settings()
    pending = any(r['status'] == 'pending' for r in rows)
    active = [r for r in rows if r['status'] in ('pending', 'confirmed')]
    paid = sum(r['paid'] for r in active)
    rate = refund_rate(settings, rows[0]['checkin'])
    can_cancel = bool(settings.get('customerCancel')) and bool(active) and rows[0]['checkin'] >= date.today().isoformat()
    return {
        'canCancel': can_cancel, 'cancelPaid': paid, 'refundRate': rate,
        'refundAmount': refund_amount(paid, rate),
        'code': rows[0]['group_code'] or rows[0]['code'], 'name': rows[0]['name'],
        'checkin': rows[0]['checkin'], 'checkout': rows[0]['checkout'], 'nights': rows[0]['nights'],
        'items': [{'room_name': names.get(r['room_id'], r['room_id']), 'guests': r['guests'], 'total': r['total'],
                   'status': r['status'], 'status_label': r['status_label'], 'options': r['options']} for r in rows],
        'total': sum(r['total'] for r in rows if r['status'] != 'cancelled'),
        'paid': sum(r['paid'] for r in rows),
        'deadline': fmt_deadline(rows[0]['deadline']) if pending else '',
        'bank': settings['bank'] if pending else '', 'phone': settings['phone'],
        'refundPolicy': refund_policy_text(settings),
    }


def refund_amount(paid, rate):
    return paid * rate // 100 // 10 * 10  # 10원 단위 절사


def customer_cancel(code, phone, data):
    """손님이 예약조회에서 직접 취소 (사용자취소). 입금한 돈이 있으면 환불 계좌 필요"""
    settings = load_settings()
    if not settings.get('customerCancel'):
        raise BookingError(f'예약 취소는 {settings["phone"]}로 연락 주세요')
    info = lookup(code, phone)
    if not info['canCancel']:
        raise BookingError(f'취소할 수 있는 예약이 없습니다. 문의 {settings["phone"]}')
    bank = {k: str(data.get(k, '')).strip()[:40] for k in ('refund_bank', 'refund_account', 'refund_holder')}
    if info['cancelPaid'] and info['refundAmount'] and not all(bank.values()):
        raise BookingError('환불받을 은행 · 계좌번호 · 예금주를 입력해 주세요')
    rate = refund_rate(settings, info['checkin'])
    with _lock:
        rows = [r for r in group_rows(info['code']) if r['status'] in ('pending', 'confirmed')
                and r['phone'] == norm_phone(phone)]
        for r in rows:
            refund = refund_amount(r['paid'], rate)
            db().execute(
                "UPDATE reservations SET status='cancelled', cancel_by='user', cancelled_at=?, deadline=NULL, "
                "cancel_reason=?, refund=?, refund_status=?, refund_bank=?, refund_account=?, refund_holder=?, "
                "updated_at=? WHERE id=?",
                (now_str(), f'고객 직접 취소 (환불 {rate}%)' if r['paid'] else '고객 직접 취소 (입금 전)', refund,
                 'due' if refund else '', bank['refund_bank'], bank['refund_account'], bank['refund_holder'],
                 now_str(), r['id']))
        db().commit()
        rows = [get_reservation(r['id']) for r in rows]
    notify('cancelled', rows)
    if settings.get('adminNotify'):
        notify('admin_cancel', rows)
    check_auto_block(phone)
    return {'refund': sum(r['refund'] for r in rows)}


def refund_done(ids, done=True, sms=True):
    """환불현황: 환불 완료 처리 (또는 환불예정으로 되돌리기). 완료 시 손님에게 환불완료 문자"""
    with _lock:
        changed = []
        for rid in ids:
            cur = db().execute("UPDATE reservations SET refund_status=?, refund_done_at=?, updated_at=? "
                               "WHERE id=? AND refund > 0 AND refund_status != ?",
                               ('done' if done else 'due', now_str() if done else None, now_str(), int(rid),
                                'done' if done else 'due'))
            if cur.rowcount:
                changed.append(int(rid))
        db().commit()
        rows = [get_reservation(r) for r in changed]
    if done and sms:
        for group in _by_group(rows):
            if group[0]['phone']:
                notify('refund_done', group)


def check_auto_block(phone):
    """예약제재: 연속 N회 취소한 고객은 자동으로 '예약불가' 분류"""
    n = int(load_settings().get('autoBlockCancels') or 0)
    phone = norm_phone(phone)
    if not n or not phone:
        return
    with _lock:
        st = _customer_stats().get(phone)
        if st and st['cancel_streak'] >= n:
            db().execute("UPDATE customers SET category=?, memo=CASE WHEN memo='' THEN ? ELSE memo END "
                         "WHERE phone=? AND category IN ('', '일반')",
                         (BLOCK_CATEGORY, f'연속 {st["cancel_streak"]}회 취소로 자동 예약제한 ({now_str()[:10]})', phone))
            db().commit()


def block_all(day, reason):
    """전체예약불가: 그날 밤 비어 있는 모든 객실을 예약막음"""
    d = parse_date(day)
    settings, prices = load_settings(), load_prices()
    made = 0
    with _lock:
        for room in prices['rooms']:
            free = units_of(settings, room['id']) - booked_counts(room['id'], d, d + timedelta(days=1)).get(d, 0)
            for _ in range(max(0, free)):
                code = new_code()
                q = {'room_id': room['id'], 'checkin': d.isoformat(), 'checkout': (d + timedelta(days=1)).isoformat(),
                     'guests': 0, 'options': [], 'room_amount': 0, 'extra_amount': 0, 'option_amount': 0}
                _insert(q, 'block', '기타', None, (reason or '전체예약불가')[:30], '', '', '', 0, 0, code, code)
                made += 1
        db().commit()
    return made


def block_range(room_id, start, end, reason):
    """휴장일 등록: 객실(또는 전체 객실)의 기간을 막음. 이미 예약된 날은 건너뛰고 빈 날만 막음"""
    s_, e_ = parse_date(start, '시작일'), parse_date(end, '종료일')
    if e_ < s_:
        raise BookingError('종료일이 시작일보다 빠릅니다')
    settings, prices = load_settings(), load_prices()
    rooms = [r for r in prices['rooms'] if not room_id or r['id'] == room_id]
    made, skipped = 0, 0
    with _lock:
        for room in rooms:
            counts = booked_counts(room['id'], s_, e_ + timedelta(days=1))
            seg = None
            d = s_
            while d <= e_ + timedelta(days=1):
                free = d <= e_ and counts.get(d, 0) < units_of(settings, room['id'])
                if d <= e_ and not free:
                    skipped += 1
                if free and seg is None:
                    seg = d
                if not free and seg is not None:
                    code = new_code()
                    q = {'room_id': room['id'], 'checkin': seg.isoformat(), 'checkout': d.isoformat(), 'guests': 0,
                         'options': [], 'room_amount': 0, 'extra_amount': 0, 'option_amount': 0}
                    _insert(q, 'block', '기타', None, (reason or '휴장')[:30], '', '', '', 0, 0, code, code)
                    made += 1
                    seg = None
                d += timedelta(days=1)
        db().commit()
    return {'made': made, 'skipped': skipped}


def monthly_report(start_month, end_month, basis='created'):
    """월별현황: 월별 총예약 · 완료 · 대기 · 환불 · 취소 · 부분취소 건수 (객실 단위)"""
    col = 'substr(created_at,1,7)' if basis == 'created' else 'substr(checkin,1,7)'
    with _lock:
        rows = [dict(r) for r in db().execute(
            f"SELECT {col} AS m, status, refund, group_code FROM reservations "
            f"WHERE status != 'block' AND {col} >= ? AND {col} <= ?", (start_month, end_month)).fetchall()]
        groups = {}
        for r in db().execute("SELECT group_code, status FROM reservations WHERE status != 'block'"):
            groups.setdefault(r['group_code'], set()).add(r['status'] == 'cancelled')
    out = {}
    for r in rows:
        m = out.setdefault(r['m'], {'month': r['m'], 'total': 0, 'confirmed': 0, 'pending': 0, 'refunded': 0,
                                    'cancelled': 0, 'partial': 0, '_groups': set()})
        m['total'] += 1
        if r['status'] in ('confirmed', 'pending'):
            m[r['status']] += 1
        if r['status'] == 'cancelled':
            m['cancelled'] += 1
            m['refunded'] += int(r['refund'] > 0)
            if groups.get(r['group_code']) == {True, False} and r['group_code'] not in m['_groups']:
                m['partial'] += 1  # 여러 객실 중 일부만 취소된 예약
                m['_groups'].add(r['group_code'])
    result = []
    for m in sorted(out):
        out[m].pop('_groups')
        result.append(out[m])
    return result


def options_calendar(month):
    """추가신청현황: 입실일별 추가상품(바비큐 등) 신청 내역"""
    first = parse_date(str(month) + '-01', '월')
    nxt = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    names = {r['id']: r['name'] for r in load_prices()['rooms']}
    days = {}
    for r in list_reservations(first.isoformat(), nxt.isoformat(), ['pending', 'confirmed'], basis='checkin'):
        for o in r['options']:
            days.setdefault(r['checkin'], []).append({
                'id': r['id'], 'name': o['name'], 'detail': o.get('detail', ''), 'qty': o['qty'],
                'guest': r['name'], 'phone': r['phone_fmt'], 'room': names.get(r['room_id'], r['room_id']),
                'status': r['status']})
    return {'month': month, 'days': days}


def bulk_sms(ids, kind, msg):
    """선택한 예약들에 문자 (같은 예약번호는 1통만)"""
    settings = load_settings()
    with _lock:
        rows = [get_reservation(int(i)) for i in ids]
    sent = 0
    for group in _by_group(rows):
        if not group[0]['phone']:
            continue
        text = render_template(settings['templates'].get(kind, '') if kind else msg, group, settings)
        if text.strip():
            threading.Thread(target=send_sms, args=(group[0]['phone'], text, kind or 'manual', group[0]['id']),
                             daemon=True).start()
            sent += 1
    return sent


# ──────────────────────────────────────────────
#  손님용 달력 / 검색
# ──────────────────────────────────────────────
def month_calendar(month):
    """해당 월 날짜별 예약 가능한 객실 목록"""
    try:
        first = date.fromisoformat(str(month) + '-01')
    except ValueError:
        raise BookingError('잘못된 월')
    nxt = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    settings, prices = load_settings(), load_prices()
    today, limit = date.today(), open_until(settings)
    holidays = set(prices.get('holidays', []))
    specials = settings.get('specialPrices') or []
    with _lock:
        expire_overdue()
        counts = {r['id']: booked_counts(r['id'], first, nxt) for r in prices['rooms']}
    days = {}
    d = first
    while d < nxt:
        open_day = today <= d <= limit
        peak, weekend = day_info(prices, d)
        days[d.isoformat()] = {
            'open': open_day, 'peak': peak, 'weekend': weekend, 'holiday': d.isoformat() in holidays,
            'rooms': [r['id'] for r in prices['rooms']
                      if open_day and counts[r['id']].get(d, 0) < units_of(settings, r['id'])],
            # 그날 밤 객실별 1박 요금 (기준인원 · 임시요금 반영) — 예약 달력에 객실 · 금액 표시용
            'prices': {r['id']: night_price(prices, r, d, specials)[0] for r in prices['rooms']} if open_day else {},
            # 할인 중인 객실: 원래 요금 · 임시요금 메모 (손님 화면에 "원래 ○원 → 할인 ○원" 표시)
            'sale': {r['id']: {'base': night_price(prices, r, d)[0],
                               'note': (find_special(r, d, specials, weekend) or {}).get('note', '')}
                     for r in prices['rooms']
                     if open_day and night_price(prices, r, d, specials)[0] < night_price(prices, r, d)[0]},
        }
        d += timedelta(days=1)
    return {'month': month, 'days': days, 'total': len(prices['rooms']),
            'roomList': [{'id': r['id'], 'name': r['name'], 'std': r['std'], 'max': r['max']} for r in prices['rooms']]}


def search(checkin, checkout):
    """선택한 날짜에 객실별 예약 가능 여부와 기준인원 요금"""
    settings, prices = load_settings(), load_prices()
    check_open_range(settings, checkin, checkout)
    check_two_night(settings, checkin, checkout)
    out = []
    with _lock:
        expire_overdue()
        for room in prices['rooms']:
            q = quote(room['id'], checkin, checkout, room['std'], [], prices, settings['maxNights'])
            out.append({
                'id': room['id'], 'name': room['name'], 'group': room['group'], 'badge': room.get('badge', ''),
                'cover': room.get('cover', ''), 'std': room['std'], 'max': room['max'],
                'extra': room.get('extra', ''), 'extraPerson': extra_person_fee(room), 'desc': room.get('desc', ''),
                'room_amount': q['room_amount'], 'nights': q['nights'],
                # 할인 전 원래 객실 요금 (같으면 할인 없음)
                'base_amount': sum(night_price(prices, room, date.fromisoformat(n['date']))[0] for n in q['nights']),
                'available': is_available(settings, room['id'], q['checkin'], q['checkout']),
            })
    return out


def public_info():
    settings, prices = load_settings(), load_prices()
    return {
        'open': bool(settings.get('bookingOpen')) and bool(settings.get('bank')),
        'phone': settings['phone'], 'depositHours': settings['depositHours'],
        'nightDeposit': settings['nightDeposit'],
        'openUntil': open_until(settings).isoformat(), 'maxNights': settings['maxNights'],
        'nextOpen': next_open_text(settings),
        'maxRooms': settings['maxRooms'], 'refundPolicy': refund_policy_text(settings),
        'customerCancel': bool(settings.get('customerCancel')),
        'notice': settings['notice'], 'today': date.today().isoformat(),
        'options': [{'i': o['i'], 'name': o['name'], 'detail': o.get('detail', ''), 'price': o['amount'],
                     'limit': (settings.get('optionLimits') or {}).get(str(o['i']), 0)}
                    for o in paid_options(prices)],
    }


# ──────────────────────────────────────────────
#  문자 (알리고)
# ──────────────────────────────────────────────
def fmt_day(ds):
    d = date.fromisoformat(ds)
    return f'{d.month}/{d.day}({WEEKDAY[d.weekday()]})'


def fmt_deadline(ts):
    if not ts:
        return ''
    t = datetime.strptime(ts, '%Y-%m-%d %H:%M:%S')
    return f'{t.month}/{t.day}({WEEKDAY[t.weekday()]}) {t.hour:02d}:{t.minute:02d}'


def render_template(tpl, rows, settings=None, drop_empty=False):
    """문자 문구의 {이름} 같은 항목을 실제 값으로 바꿈. rows = 같은 예약번호로 묶인 객실들
    drop_empty=True 면 값이 빈 항목이 들어 있는 줄은 통째로 뺌 (텔레그램 알림용)"""
    settings = settings or load_settings()
    names = {r['id']: r['name'] for r in load_prices()['rooms']}
    r = rows[0]
    opt_text = lambda x: ', '.join(f'{o["name"]} {o.get("detail", "")} x{o["qty"]}'.replace('  ', ' ') for o in x['options'])
    vals = {
        '객실상세': '\n'.join(f'· {names.get(x["room_id"], x["room_id"])} {x["guests"]}명 {x["total"]:,}원'
                          + (f' ({opt_text(x)})' if opt_text(x) else '') for x in rows),
        '옵션': ', '.join(t for t in (opt_text(x) for x in rows) if t),
        '요청사항': r.get('request') or '', '메모': r.get('memo') or '', '예약경로': r.get('source') or '',
        '환불계좌': r.get('refund_info') or '',
        '예약번호': r['group_code'] or r['code'], '이름': r['name'], '연락처': fmt_phone(r['phone']),
        '객실': ', '.join(names.get(x['room_id'], x['room_id']) for x in rows),
        '입실': fmt_day(r['checkin']), '퇴실': fmt_day(r['checkout']), '박': r['nights'],
        '인원': sum(x['guests'] for x in rows),
        '금액': f"{sum(x['total'] for x in rows):,}", '잔금': f"{sum(x['total'] - x['paid'] for x in rows):,}",
        '환불': f"{sum(x['refund'] for x in rows):,}",
        '입금액': f"{sum(x['paid'] for x in rows):,}", '결제방법': r.get('pay_method') or '',
        '계좌': settings.get('bankShort') or settings['bank'], '입금기한': fmt_deadline(r.get('deadline')),
        '펜션전화': settings['phone'],
    }
    sub = lambda text: re.sub(r'\{([^{}]+)\}', lambda m: str(vals.get(m.group(1), m.group(0))), text)
    if not drop_empty:
        return sub(tpl)
    keep = [line for line in tpl.split('\n')
            if all(str(vals.get(k, 'x')).strip() not in ('', '0') for k in re.findall(r'\{([^{}]+)\}', line))]
    return sub('\n'.join(keep))


def aligo_call(path, fields, settings):
    a = settings['aligo']
    body = urllib.parse.urlencode({'key': a['key'], 'user_id': a['userId'], **fields}).encode('utf-8')
    req = urllib.request.Request('https://apis.aligo.in/' + path, data=body)
    with urllib.request.urlopen(req, timeout=15) as resp:
        return json.loads(resp.read().decode('utf-8'))


def send_sms(phone, msg, kind='manual', reservation_id=None):
    """문자 1건 발송 후 기록. (성공 여부, 결과 메시지)"""
    settings = load_settings()
    a = settings['aligo']
    phone = norm_phone(phone)
    if not valid_mobile(phone):
        ok, result = False, '휴대폰 번호 없음/오류'
    elif not (a.get('enabled') and a.get('key') and a.get('userId') and a.get('sender')):
        ok, result = False, '문자 발송 꺼짐 (설정에서 알리고 연동 필요) — 기록만 저장'
    else:
        lms = len(msg.encode('euc-kr', errors='replace')) > 90  # 90바이트 넘으면 장문(LMS)
        fields = {'sender': norm_phone(a['sender']), 'receiver': phone, 'msg': msg,
                  'msg_type': 'LMS' if lms else 'SMS'}
        if lms:
            fields['title'] = '동해카라반펜션'
        if a.get('testmode'):
            fields['testmode_yn'] = 'Y'
        try:
            res = aligo_call('send/', fields, settings)
            ok = int(res.get('result_code', -1)) > 0
            result = str(res.get('message', '')) + (' · 장문' if lms else '') + \
                (' (테스트모드 — 실제 발송 안 됨)' if a.get('testmode') else '')
        except Exception as e:  # 네트워크 오류 등
            ok, result = False, f'발송 오류: {e}'
    with _lock:
        db().execute('INSERT INTO sms_log (created_at, reservation_id, phone, kind, msg, ok, result) '
                     'VALUES (?,?,?,?,?,?,?)', (now_str(), reservation_id, phone, kind, msg, int(ok), result))
        db().commit()
    return ok, result


TELEGRAM_LOG = '텔레그램'  # 발송 기록의 받는 번호 칸 (문자비 계산에서 제외)


def telegram_call(method, params, settings=None):
    settings = settings or load_settings()
    token = (settings.get('telegram') or {}).get('botToken', '').strip()
    if not token:
        raise BookingError('텔레그램 봇 토큰을 먼저 입력하세요')
    body = urllib.parse.urlencode(params).encode('utf-8')
    req = urllib.request.Request(f'https://api.telegram.org/bot{token}/{method}', data=body)
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            return json.loads(resp.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        try:
            msg = json.loads(e.read().decode('utf-8')).get('description', str(e))
        except Exception:
            msg = str(e)
        raise BookingError(f'텔레그램 오류: {msg}')
    except Exception as e:
        raise BookingError(f'텔레그램 연결 실패: {e}')


def send_telegram(text, kind='admin', reservation_id=None):
    settings = load_settings()
    tg = settings.get('telegram') or {}
    try:
        if not tg.get('chatId'):
            raise BookingError('텔레그램 채팅 ID가 없습니다')
        res = telegram_call('sendMessage', {'chat_id': tg['chatId'], 'text': text}, settings)
        ok, result = bool(res.get('ok')), '텔레그램 (무료)'
    except BookingError as e:
        ok, result = False, str(e)
    with _lock:
        db().execute('INSERT INTO sms_log (created_at, reservation_id, phone, kind, msg, ok, result) '
                     'VALUES (?,?,?,?,?,?,?)', (now_str(), reservation_id, TELEGRAM_LOG, kind, text, int(ok), result))
        db().commit()
    return ok, result


def telegram_find_chat():
    """봇에게 아무 메시지나 보낸 뒤 누르면 채팅 ID 를 찾아줌"""
    res = telegram_call('getUpdates', {'limit': 20})
    for u in reversed(res.get('result', [])):
        msg = u.get('message') or u.get('channel_post') or {}
        chat = msg.get('chat')
        if chat:
            name = chat.get('title') or ' '.join(x for x in (chat.get('last_name'), chat.get('first_name')) if x)
            return {'chatId': str(chat['id']), 'name': name or chat.get('username', '')}
    raise BookingError('봇에게 온 메시지가 없습니다. 텔레그램에서 봇에게 아무 말이나 보낸 뒤 다시 눌러 주세요')


def mask_name(name):
    """홍길동 → 홍*동, 김철 → 김*"""
    name = str(name or '')
    if len(name) <= 1:
        return name
    return name[0] + '*' * max(1, len(name) - 2) + (name[-1] if len(name) > 2 else '')


def admin_detail_text(kind, rows, settings):
    """사장님 텔레그램 알림 — 운영관리에서 고친 문구(tgTemplates) 사용, 비어 있으면 기본 문구"""
    tpl = (settings.get('tgTemplates') or {}).get(kind) or DEFAULT_SETTINGS['tgTemplates'].get(kind)
    if tpl:
        return render_template(tpl, rows, settings, drop_empty=True)
    names = {r['id']: r['name'] for r in load_prices()['rooms']}
    r = rows[0]
    title = {'admin': '🔔 새 예약 (입금대기)', 'admin_cancel': '↩️ 손님 직접 취소',
             'admin_expired': '⏰ 미입금 자동취소', 'admin_paid': '💰 입금 확인 (예약확정)'}.get(kind, '알림')
    lines = [title, f'{mask_name(r["name"])} (끝번호 {r["phone"][-4:] or "-"})', f'{fmt_day(r["checkin"])} ~ {fmt_day(r["checkout"])} ({r["nights"]}박)']
    for x in rows:
        opts = ', '.join(f'{o["name"]} {o.get("detail", "")} x{o["qty"]}' for o in x['options'])
        lines.append(f'· {names.get(x["room_id"], x["room_id"])} {x["guests"]}명 {x["total"]:,}원' + (f' ({opts})' if opts else ''))
    lines.append(f'합계 {sum(x["total"] for x in rows):,}원')
    if kind == 'admin' and r.get('deadline'):
        lines.append(f'입금기한 {fmt_deadline(r["deadline"])}')
    if kind == 'admin_cancel':
        refund = sum(x['refund'] for x in rows)
        lines.append(f'환불할 금액 {refund:,}원' + (' (환불 계좌는 관리자 화면에서 확인)' if refund else ''))
    if r.get('request'):
        lines.append('요청사항 있음 (관리자 화면에서 확인)')
    lines.append(f'예약번호 {r["group_code"] or r["code"]}')
    return '\n'.join(lines)


def notify(kind, rows, force=False):
    """자동 문자 (백그라운드로 발송 — 예약 처리가 문자 때문에 느려지지 않게)
    '문자 안 보냄' 예약은 손님 자동 문자를 건너뜀. force=True 는 관리자가 화면에서 직접 '문자 보내기'를 고른 경우"""
    settings = load_settings()
    if not kind.startswith('admin') and not force and rows[0].get('no_sms'):
        return
    if kind.startswith('admin'):
        tg = settings.get('telegram') or {}
        if tg.get('enabled') and tg.get('chatId'):  # 사장님 알림은 텔레그램 (무료, 자세히)
            threading.Thread(target=send_telegram, args=(admin_detail_text(kind, rows, settings), kind, rows[0]['id']),
                             daemon=True).start()
            return
        if kind in ('admin_expired', 'admin_paid'):  # 자동취소 · 입금확인 알림은 텔레그램일 때만 (문자비 절약)
            return
    elif not settings['auto'].get(kind, True):
        return
    tpl = settings['templates'].get(kind, '')
    if not tpl.strip():
        return
    msg = render_template(tpl, rows, settings)
    phone = settings['adminPhone'] if kind.startswith('admin') else rows[0]['phone']
    threading.Thread(target=send_sms, args=(phone, msg, kind, rows[0]['id']), daemon=True).start()


def sms_balance():
    settings = load_settings()
    a = settings['aligo']
    if not (a.get('key') and a.get('userId')):
        raise BookingError('알리고 API 키와 아이디를 먼저 입력하세요')
    try:
        res = aligo_call('remain/', {}, settings)
    except Exception as e:
        raise BookingError(f'알리고 연결 실패: {e}')
    if int(res.get('result_code', -1)) <= 0:
        raise BookingError(f'알리고 오류: {res.get("message")}')
    return {'sms': res.get('SMS_CNT'), 'lms': res.get('LMS_CNT'), 'mms': res.get('MMS_CNT')}


def sms_log(limit=300):
    with _lock:
        return [dict(r) for r in db().execute('SELECT * FROM sms_log ORDER BY id DESC LIMIT ?', (limit,)).fetchall()]


# ──────────────────────────────────────────────
#  자동 처리 (입금기한 초과 취소, 입실 전날 안내)
# ──────────────────────────────────────────────
def _by_group(rows):
    groups = {}
    for r in rows:
        groups.setdefault(r['group_code'] or r['code'], []).append(r)
    return list(groups.values())


def expire_overdue():
    with _lock:
        ids = [r['id'] for r in db().execute(
            "SELECT id FROM reservations WHERE status='pending' AND deadline IS NOT NULL AND deadline < ?",
            (now_str(),)).fetchall()]
        if not ids:
            return
        for rid in ids:
            db().execute("UPDATE reservations SET status='cancelled', cancel_reason='입금기한 초과 자동취소', "
                         "cancel_by='auto', cancelled_at=?, updated_at=? WHERE id=?", (now_str(), now_str(), rid))
        db().commit()
        rows = [get_reservation(rid) for rid in ids]
    for group in _by_group(rows):
        if group[0]['phone']:
            notify('expired', group)
            check_auto_block(group[0]['phone'])
        notify('admin_expired', group)


def send_deposit_reminders():
    settings = load_settings()
    minutes = settings.get('depositReminderMinutes') or 0
    if not minutes:
        return
    soon = (datetime.now() + timedelta(minutes=minutes)).strftime('%Y-%m-%d %H:%M:%S')
    with _lock:
        rows = [row_dict(r) for r in db().execute(
            "SELECT * FROM reservations WHERE status='pending' AND deposit_reminded=0 AND deadline IS NOT NULL "
            "AND deadline <= ? AND deadline > ? AND created_at <= ?",
            (soon, now_str(), (datetime.now() - timedelta(minutes=minutes)).strftime('%Y-%m-%d %H:%M:%S'))).fetchall()]
        for r in rows:
            db().execute('UPDATE reservations SET deposit_reminded=1 WHERE id=?', (r['id'],))
        db().commit()
    for group in _by_group(rows):
        if group[0]['phone']:
            notify('deposit_reminder', group)


def send_reminders():
    settings = load_settings()
    if datetime.now().hour < settings['reminderHour']:
        return
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    with _lock:
        rows = [row_dict(r) for r in db().execute(
            "SELECT * FROM reservations WHERE status='confirmed' AND checkin=? AND reminded=0", (tomorrow,)).fetchall()]
        for r in rows:
            db().execute('UPDATE reservations SET reminded=1 WHERE id=?', (r['id'],))
        db().commit()
    for group in _by_group(rows):
        if group[0]['phone']:
            notify('reminder', group)


def daily_backup(keep=60):
    """private/backup/ 에 하루 한 번 예약 DB 복사 (최근 60일치 보관)"""
    folder = os.path.join(PRIVATE_DIR, 'backup')
    target = os.path.join(folder, f'booking-{date.today().isoformat()}.db')
    if os.path.exists(target):
        return
    os.makedirs(folder, exist_ok=True)
    with _lock:
        dst = sqlite3.connect(target)
        db().backup(dst)
        dst.close()
    for old in sorted(f for f in os.listdir(folder) if f.startswith('booking-'))[:-keep]:
        os.remove(os.path.join(folder, old))


def backup_zip():
    """지금 시점 예약 DB + 예약 설정을 zip 하나로 (관리자 화면 [백업 내려받기])"""
    import zipfile
    tmp = os.path.join(PRIVATE_DIR, '_download.db')
    with _lock:
        dst = sqlite3.connect(tmp)
        db().backup(dst)
        dst.close()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(tmp, 'booking.db')
        if os.path.exists(SETTINGS_FILE):
            z.write(SETTINGS_FILE, 'booking_settings.json')
    os.remove(tmp)
    return buf.getvalue()


def purge_old_personal_info():
    """보관 기간(KEEP_YEARS)이 지난 예약의 개인정보 파기 — 하루 한 번
    예약 기록은 매출 통계용으로 날짜 · 객실 · 금액만 남기고 이름 · 연락처 · 요청 · 메모 · 환불계좌를 지움"""
    cut = (date.today() - timedelta(days=366 * KEEP_YEARS)).isoformat()
    with _lock:
        n = db().execute(
            "UPDATE reservations SET name='(파기)', phone='', request='', memo='', refund_bank='', refund_account='', "
            "refund_holder='', updated_at=? WHERE checkout < ? AND status != 'block' AND name != '(파기)'",
            (now_str(), cut)).rowcount
        db().execute('DELETE FROM sms_log WHERE created_at < ?', (cut,))
        # 최근 보관 기간 안에 예약이 없는 고객 명단도 삭제
        db().execute("DELETE FROM customers WHERE created_at < ? AND phone NOT IN "
                     "(SELECT phone FROM reservations WHERE phone != '')", (cut,))
        db().commit()
    if n:
        print(f'[개인정보 파기] {KEEP_YEARS}년 지난 예약 {n}건')


_last_purge = None


def _worker():
    global _last_purge
    while True:
        try:
            expire_overdue()
            send_deposit_reminders()
            send_reminders()
            daily_backup()
            if _last_purge != date.today():
                purge_old_personal_info()
                _last_purge = date.today()
        except Exception as e:
            print('[예약 자동처리 오류]', e)
        time.sleep(60)


def start_worker():
    with _lock:
        db()
    threading.Thread(target=_worker, daemon=True).start()


# ──────────────────────────────────────────────
#  통계 / 엑셀
# ──────────────────────────────────────────────
def revenue_of(r):
    """매출: 확정 예약은 총액, 취소 예약은 받은 돈 - 환불액 (위약금)"""
    if r['status'] == 'confirmed':
        return r['total']
    if r['status'] == 'cancelled':
        return max(0, r['paid'] - r['refund'])
    return 0


def month_days(year, month):
    first = date(year, month, 1)
    return ((first.replace(day=28) + timedelta(days=4)).replace(day=1) - first).days


def stats(year):
    prices, settings = load_prices(), load_settings()
    rooms = [{'id': r['id'], 'name': r['name']} for r in prices['rooms']]
    with _lock:
        rows = [row_dict(r) for r in db().execute(
            "SELECT * FROM reservations WHERE status != 'block' AND checkin < ? AND checkout > ?",
            (f'{year + 1}-01-01', f'{year}-01-01')).fetchall()]
    months = [{'month': m, 'revenue': 0, 'count': 0, 'nights': 0, 'cancelled': 0} for m in range(1, 13)]
    by_room = {r['id']: {'nights': 0, 'revenue': 0, 'count': 0} for r in rooms}
    by_source = {}
    for r in rows:
        ci = date.fromisoformat(r['checkin'])
        if ci.year == year:  # 매출 · 건수는 입실월 기준
            m = months[ci.month - 1]
            if r['status'] == 'cancelled':
                m['cancelled'] += 1
            elif r['status'] == 'confirmed':
                m['count'] += 1
                by_source[r['source']] = by_source.get(r['source'], 0) + 1
            m['revenue'] += revenue_of(r)
            if r['room_id'] in by_room:
                by_room[r['room_id']]['revenue'] += revenue_of(r)
                by_room[r['room_id']]['count'] += int(r['status'] == 'confirmed')
        if r['status'] == 'confirmed':  # 가동 박수는 실제 숙박한 날짜 기준
            d, end = ci, date.fromisoformat(r['checkout'])
            while d < end:
                if d.year == year:
                    months[d.month - 1]['nights'] += 1
                    if r['room_id'] in by_room:
                        by_room[r['room_id']]['nights'] += 1
                d += timedelta(days=1)
    total_units = sum(units_of(settings, r['id']) for r in rooms)
    for m in months:
        cap = month_days(year, m['month']) * total_units
        m['occupancy'] = round(m['nights'] / cap * 100, 1) if cap else 0
    year_days = (date(year + 1, 1, 1) - date(year, 1, 1)).days
    for r in rooms:
        s = by_room[r['id']]
        r.update(s, occupancy=round(s['nights'] / (year_days * units_of(settings, r['id'])) * 100, 1))
    return {'year': year, 'months': months, 'rooms': rooms, 'sources': by_source,
            'pending': sum(1 for r in rows if r['status'] == 'pending'),
            'revenue': sum(m['revenue'] for m in months)}


EXPORT_COLS = [('예약번호', 'code'), ('상태', 'status_label'), ('경로', 'source'), ('객실', 'room_name'),
               ('입실', 'checkin'), ('퇴실', 'checkout'), ('박', 'nights'), ('인원', 'guests'),
               ('예약자', 'name'), ('연락처', 'phone_fmt'), ('객실요금', 'room_amount'), ('인원추가', 'extra_amount'),
               ('옵션', 'option_text'), ('옵션요금', 'option_amount'), ('총액', 'total'), ('입금액', 'paid'),
               ('환불액', 'refund'), ('결제구분', 'pay_method'), ('취소구분', 'cancel_by_label'),
               ('환불상태', 'refund_label'), ('환불계좌', 'refund_info'), ('요청사항', 'request'), ('메모', 'memo'), ('취소사유', 'cancel_reason'),
               ('접수일시', 'created_at')]
EXPORT_WIDTHS = (16, 9, 9, 13, 11, 11, 5, 5, 10, 14, 10, 9, 26, 9, 10, 10, 9, 9, 11, 9, 28, 30, 30, 20, 19)
MONEY_COLS = {'room_amount', 'extra_amount', 'option_amount', 'total', 'paid', 'refund'}


def export(start, end, **filters):
    """엑셀(.xlsx) 파일 내용 반환. openpyxl 이 없으면 CSV"""
    names = {r['id']: r['name'] for r in load_prices()['rooms']}
    rows = list_reservations(start, end, **filters)
    for r in rows:
        r['room_name'] = names.get(r['room_id'], r['room_id'])
        r['option_text'] = ', '.join(f"{o['name']}({o.get('detail', '')}) x{o['qty']}" for o in r['options'])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        import csv
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([c[0] for c in EXPORT_COLS])
        for r in rows:
            w.writerow([r.get(k, '') for _, k in EXPORT_COLS])
        return ('﻿' + buf.getvalue()).encode('utf-8'), 'text/csv; charset=utf-8', 'csv'
    wb = Workbook()
    ws = wb.active
    ws.title = '예약목록'
    ws.append([c[0] for c in EXPORT_COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0E6BA8')
    for r in rows:
        ws.append([r.get(k, '') for _, k in EXPORT_COLS])
    for i, (w, (_, key)) in enumerate(zip(EXPORT_WIDTHS, EXPORT_COLS), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        if key in MONEY_COLS:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = '#,##0'
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'


# ──────────────────────────────────────────────
#  고객관리 (땡큐캠핑 캠퍼관리)
# ──────────────────────────────────────────────
BLOCK_CATEGORY = '예약불가'


def is_blocked_customer(phone):
    with _lock:
        r = db().execute('SELECT category, deleted FROM customers WHERE phone=?', (norm_phone(phone),)).fetchone()
    return bool(r and r['category'] == BLOCK_CATEGORY and not r['deleted'])


def _customer_stats():
    """전화번호별 누적 예약 건수 · 금액 · 연속 취소 횟수"""
    stats = {}
    rows = db().execute("SELECT phone, group_code, status, total, paid, refund, created_at FROM reservations "
                        "WHERE phone != '' AND status != 'block' ORDER BY created_at DESC").fetchall()
    seen = {}
    for r in rows:
        st = stats.setdefault(r['phone'], {'count': 0, 'amount': 0, 'cancel_streak': 0, '_streak_open': True,
                                           'last': ''})
        st['last'] = st['last'] or r['created_at']
        if r['status'] == 'confirmed':
            st['count'] += 1
            st['amount'] += r['total']
        g = (r['phone'], r['group_code'])
        if g in seen:
            continue
        seen[g] = True
        if st['_streak_open']:  # 최근 예약부터 연속으로 취소된 횟수 (예약번호 단위)
            if r['status'] == 'cancelled':
                st['cancel_streak'] += 1
            else:
                st['_streak_open'] = False
    for st in stats.values():
        st.pop('_streak_open')
    return stats


def list_customers(view='all', q='', category=''):
    sql, args = 'SELECT * FROM customers WHERE deleted=?', [1 if view == 'deleted' else 0]
    if view == 'blocked':
        sql += ' AND category=?'
        args.append(BLOCK_CATEGORY)
    if category:
        sql += ' AND category=?'
        args.append(category)
    if q:
        sql += ' AND (name LIKE ? OR nickname LIKE ? OR memo LIKE ?' + (' OR phone LIKE ?' if norm_phone(q) else '') + ')'
        args += [f'%{q}%'] * 3 + ([f'%{norm_phone(q)}%'] if norm_phone(q) else [])
    sql += ' ORDER BY created_at DESC LIMIT 5000'
    with _lock:
        rows = [dict(r) for r in db().execute(sql, args).fetchall()]
        stats = _customer_stats()
        counts = {
            'all': db().execute('SELECT COUNT(*) FROM customers WHERE deleted=0').fetchone()[0],
            'blocked': db().execute('SELECT COUNT(*) FROM customers WHERE deleted=0 AND category=?',
                                    (BLOCK_CATEGORY,)).fetchone()[0],
            'deleted': db().execute('SELECT COUNT(*) FROM customers WHERE deleted=1').fetchone()[0],
        }
    empty = {'count': 0, 'amount': 0, 'cancel_streak': 0, 'last': ''}
    for r in rows:
        r.update(stats.get(r['phone'], empty))
        r['phone_fmt'] = fmt_phone(r['phone'])
    return {'items': rows, 'counts': counts}


def get_customer(phone):
    phone = norm_phone(phone)
    with _lock:
        c = db().execute('SELECT * FROM customers WHERE phone=?', (phone,)).fetchone()
        if not c:
            raise BookingError('고객을 찾을 수 없습니다')
        res = [row_dict(r) for r in db().execute(
            'SELECT * FROM reservations WHERE phone=? ORDER BY checkin DESC', (phone,)).fetchall()]
        c = dict(c)
        c.update(_customer_stats().get(phone, {'count': 0, 'amount': 0, 'cancel_streak': 0, 'last': ''}))
    c['phone_fmt'] = fmt_phone(phone)
    return {'item': c, 'reservations': res}


def save_customer(data):
    """고객 등록(새 번호) 또는 수정"""
    phone = norm_phone(data.get('phone'))
    if not valid_mobile(phone):
        raise BookingError('휴대폰 번호를 정확히 입력해 주세요')
    fields = {k: str(data.get(k, '')).strip()[:200] for k in ('name', 'nickname', 'category', 'memo')}
    with _lock:
        db().execute('INSERT OR IGNORE INTO customers (phone, name, created_at) VALUES (?,?,?)',
                     (phone, fields['name'], now_str()))
        db().execute('UPDATE customers SET name=?, nickname=?, category=?, memo=?, deleted=0 WHERE phone=?',
                     (fields['name'], fields['nickname'], fields['category'], fields['memo'], phone))
        db().commit()
    return get_customer(phone)


def customers_action(phones, action, category=''):
    """선택한 고객: 분류 이동 / 삭제 / 복구"""
    phones = [norm_phone(p) for p in phones if norm_phone(p)]
    with _lock:
        for p in phones:
            if action == 'category':
                db().execute('UPDATE customers SET category=? WHERE phone=?', (category, p))
            elif action == 'delete':
                db().execute('UPDATE customers SET deleted=1 WHERE phone=?', (p,))
            elif action == 'restore':
                db().execute('UPDATE customers SET deleted=0 WHERE phone=?', (p,))
            elif action == 'purge':  # 삭제고객에서 완전 삭제
                db().execute('DELETE FROM customers WHERE phone=? AND deleted=1', (p,))
        db().commit()
    return len(phones)


def import_customers(data, filename):
    """땡큐캠핑 등에서 받은 엑셀(.xlsx) / CSV 파일에서 이름 · 휴대폰 가져오기"""
    rows = []
    if filename.lower().endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise BookingError('엑셀 파일을 읽을 수 없습니다. CSV 로 저장해서 올려 주세요')
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [['' if v is None else str(v) for v in r] for r in wb.active.iter_rows(values_only=True)]
    else:
        import csv
        for enc in ('utf-8-sig', 'cp949'):
            try:
                rows = list(csv.reader(io.StringIO(data.decode(enc))))
                break
            except UnicodeDecodeError:
                continue
    # 머리글 줄에서 이름 · 휴대폰 칸 찾기
    head_i = name_c = phone_c = nick_c = None
    for i, r in enumerate(rows[:20]):
        cells = [c.replace(' ', '') for c in r]
        p = next((k for k, c in enumerate(cells) if c in ('휴대폰', '휴대폰번호', '연락처', '전화번호', '핸드폰')), None)
        n = next((k for k, c in enumerate(cells) if c in ('고객명', '이름', '예약자', '예약자명', '성명')), None)
        if p is not None:
            head_i, phone_c, name_c = i, p, n
            nick_c = next((k for k, c in enumerate(cells) if c == '닉네임'), None)
            break
    if head_i is None:
        raise BookingError('파일 첫 줄에 "고객명", "휴대폰" 제목 칸이 있어야 합니다')
    added = updated = skipped = 0
    with _lock:
        for r in rows[head_i + 1:]:
            phone = norm_phone(r[phone_c] if phone_c < len(r) else '')
            if not valid_mobile(phone):
                skipped += 1
                continue
            name = (r[name_c] if name_c is not None and name_c < len(r) else '').strip()[:30]
            nick = (r[nick_c] if nick_c is not None and nick_c < len(r) else '').strip()[:30]
            cur = db().execute('INSERT OR IGNORE INTO customers (phone, name, nickname, created_at) VALUES (?,?,?,?)',
                               (phone, name, nick, now_str()))
            if cur.rowcount:
                added += 1
            else:
                updated += 1
                if name:
                    db().execute("UPDATE customers SET name=CASE WHEN name='' THEN ? ELSE name END WHERE phone=?",
                                 (name, phone))
        db().commit()
    return {'added': added, 'existing': updated, 'skipped': skipped}


def customers_sms(phones, msg, category_all=None):
    """고객 문자 발송 (선택 고객 또는 분류 전체). 순서대로 백그라운드 발송"""
    msg = str(msg or '').strip()
    if not msg:
        raise BookingError('문자 내용을 입력하세요')
    if category_all is not None:
        with _lock:
            sql = "SELECT phone FROM customers WHERE deleted=0 AND category != ?"
            args = [BLOCK_CATEGORY]
            if category_all:
                sql += ' AND category=?'
                args.append(category_all)
            phones = [r['phone'] for r in db().execute(sql, args).fetchall()]
    phones = sorted({norm_phone(p) for p in phones if valid_mobile(p)})

    def run():
        for p in phones:
            send_sms(p, msg, 'customer')
    threading.Thread(target=run, daemon=True).start()
    return len(phones)


# ──────────────────────────────────────────────
#  매출관리 (땡큐캠핑 매출관리와 같은 기준: 퇴실일자 기준)
# ──────────────────────────────────────────────
def _sales_rows(start, end, basis='checkout', exclude_admin=False, pay_method=None):
    """매출 대상 예약 (확정 + 위약금이 남은 취소). start~end(미포함) 기준일 범위"""
    col = {'checkout': 'checkout', 'checkin': 'checkin'}.get(basis, 'checkout')
    sql = (f"SELECT * FROM reservations WHERE status IN ('confirmed','cancelled') AND {col} >= ? AND {col} < ?")
    args = [start, end]
    if exclude_admin:
        sql += " AND source = '홈페이지'"
    if pay_method:
        sql += ' AND pay_method = ?'
        args.append(pay_method)
    with _lock:
        rows = [row_dict(r) for r in db().execute(sql, args).fetchall()]
    for r in rows:
        r['_day'] = r[col]
    return rows


def _sales_bucket(rows, key):
    out = {}
    for r in rows:
        b = out.setdefault(key(r), {'count': 0, 'amount': 0, 'deduct_count': 0, 'deduct_amount': 0})
        if r['status'] == 'confirmed':
            b['count'] += 1
            b['amount'] += r['total']
        elif revenue_of(r) > 0:  # 취소됐지만 위약금(받은 돈 - 환불)이 남은 건
            b['deduct_count'] += 1
            b['deduct_amount'] += revenue_of(r)
    for b in out.values():
        b['total'] = b['amount'] + b['deduct_amount']
    return out


def sales_monthly(year, basis='checkout', exclude_admin=False, pay_method=None):
    rows = _sales_rows(f'{year}-01-01', f'{year + 1}-01-01', basis, exclude_admin, pay_method)
    b = _sales_bucket(rows, lambda r: int(r['_day'][5:7]))
    empty = {'count': 0, 'amount': 0, 'deduct_count': 0, 'deduct_amount': 0, 'total': 0}
    return [dict(empty, **b.get(m, {}), month=m) for m in range(1, 13)]


def sales_daily(year, month, basis='checkout', exclude_admin=False, pay_method=None):
    first = date(year, month, 1)
    nxt = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    rows = _sales_rows(first.isoformat(), nxt.isoformat(), basis, exclude_admin, pay_method)
    b = _sales_bucket(rows, lambda r: r['_day'])
    empty = {'count': 0, 'amount': 0, 'deduct_count': 0, 'deduct_amount': 0, 'total': 0}
    out, d = [], first
    while d < nxt:
        out.append(dict(empty, **b.get(d.isoformat(), {}), date=d.isoformat()))
        d += timedelta(days=1)
    return out


def money_calendar(month):
    """예약정산달력: 날짜별 실제 입금 · 환불 금액"""
    first = parse_date(str(month) + '-01', '월')
    nxt = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    days = {}
    with _lock:
        for r in db().execute("SELECT substr(paid_at,1,10) d, COUNT(*) n, SUM(paid) s FROM reservations "
                              "WHERE paid > 0 AND paid_at >= ? AND paid_at < ? GROUP BY d",
                              (first.isoformat(), nxt.isoformat())):
            days.setdefault(r['d'], {})['paid'] = {'count': r['n'], 'amount': r['s']}
        for r in db().execute("SELECT substr(refund_done_at,1,10) d, COUNT(*) n, SUM(refund) s FROM reservations "
                              "WHERE refund_status='done' AND refund_done_at >= ? AND refund_done_at < ? GROUP BY d",
                              (first.isoformat(), nxt.isoformat())):
            days.setdefault(r['d'], {})['refund'] = {'count': r['n'], 'amount': r['s']}
        due = db().execute("SELECT COUNT(*) n, COALESCE(SUM(refund),0) s FROM reservations "
                           "WHERE refund_status='due'").fetchone()
    return {'month': month, 'days': days, 'refundDue': {'count': due['n'], 'amount': due['s']}}


def cost_report(year):
    """비용 · 정산자료: 월별 매출 - 문자비 - 고정비 = 순매출 (+ 땡큐캠핑 수수료와 비교)"""
    settings = load_settings()
    months = sales_monthly(year)
    sms = {m: {'sms': 0, 'lms': 0} for m in range(1, 13)}
    with _lock:
        for r in db().execute("SELECT created_at, msg FROM sms_log WHERE ok=1 AND phone != ? "
                              "AND created_at >= ? AND created_at < ?",
                              (TELEGRAM_LOG, f'{year}-01-01', f'{year + 1}-01-01')):
            kind = 'lms' if len(r['msg'].encode('euc-kr', errors='replace')) > 90 else 'sms'
            sms[int(r['created_at'][5:7])][kind] += 1
    fixed = int(settings.get('monthlyFixedCost') or 0)
    rate = float(settings.get('compareFeeRate') or 0)
    out = []
    for m in months:
        c = sms[m['month']]
        sms_cost = round(c['sms'] * float(settings.get('smsPrice') or 0) + c['lms'] * float(settings.get('lmsPrice') or 0))
        has_data = m['total'] or c['sms'] or c['lms']
        fixed_m = fixed if has_data else 0
        out.append({'month': m['month'], 'count': m['count'] + m['deduct_count'], 'revenue': m['total'],
                    'sms': c['sms'], 'lms': c['lms'], 'sms_cost': sms_cost, 'fixed': fixed_m,
                    'net': m['total'] - sms_cost - fixed_m,
                    'compare_fee': round(m['total'] * rate / 100),
                    'saved': round(m['total'] * rate / 100) - sms_cost - fixed_m})
    return {'year': year, 'months': out, 'rate': rate, 'smsPrice': settings.get('smsPrice'),
            'lmsPrice': settings.get('lmsPrice'), 'fixed': fixed}


# ──────────────────────────────────────────────
#  요금표의 성수기 기간 · 공휴일 수정
# ──────────────────────────────────────────────
def save_calendar_rules(peak_periods, holidays):
    periods = []
    for p in peak_periods:
        s, e = parse_date(p.get('start'), '성수기 시작일'), parse_date(p.get('end'), '성수기 종료일')
        if e < s:
            raise BookingError('성수기 종료일이 시작일보다 빠릅니다')
        periods.append({'start': s.isoformat(), 'end': e.isoformat()})
    hol = sorted({parse_date(h, '공휴일').isoformat() for h in holidays if str(h).strip()})
    with _lock:
        prices = load_prices()
        prices['peakPeriods'] = sorted(periods, key=lambda p: p['start'])
        prices['holidays'] = hol
        with open(PRICES_FILE, 'w', encoding='utf-8') as f:
            json.dump(prices, f, ensure_ascii=False, indent=2)
    return prices


# ──────────────────────────────────────────────
#  HTTP 요청 처리 (server.py 가 호출)
# ──────────────────────────────────────────────
_hits = {}  # 예약 시도 횟수 제한 (IP별)


def rate_limited(key, limit, window):
    t = time.time()
    with _lock:
        hits = [h for h in _hits.get(key, []) if t - h < window]
        if len(hits) >= limit:
            _hits[key] = hits
            return True
        hits.append(t)
        _hits[key] = hits
        return False


# ──────────────────────────────────────────────
#  홈페이지 방문자 수 (관리자 [운영관리 → 방문자 통계])
#  손님 브라우저가 만든 임의 번호(vid)로 "오늘 처음 온 사람"만 셉니다.
#  이름·전화·IP 는 저장하지 않고, 임의 번호도 이틀 뒤 지우고 날짜별 숫자만 남깁니다.
# ──────────────────────────────────────────────
VISIT_PAGES = ('home', 'reserve')
VISIT_ID_RE = re.compile(r'^[0-9a-zA-Z]{8,40}$')
BOT_RE = re.compile(r'bot|crawl|spider|slurp|preview|headless|lighthouse|facebookexternalhit', re.I)
VISIT_SOURCES = [  # 어디서 들어왔는지 (이전 페이지 주소로 판단)
    ('네이버', ('naver.',)), ('구글', ('google.',)), ('다음·카카오', ('daum.', 'kakao.')),
    ('인스타그램', ('instagram.',)), ('페이스북', ('facebook.', 'fb.')), ('유튜브', ('youtube.', 'youtu.be')),
    ('땡큐캠핑', ('thankqcamping.',)),
]


def visit_source(ref):
    host = re.sub(r'^https?://', '', (ref or '').strip().lower()).split('/')[0]
    if not host or 'syyoon2-dhc.github.io' in host or host.endswith('.ts.net') or host.startswith(('localhost', '127.')):
        return '직접 방문'
    for name, keys in VISIT_SOURCES:
        if any(k in host for k in keys):
            return name
    return '기타 사이트'


def record_visit(body):
    page, vid, ua = body.get('page'), str(body.get('vid') or ''), str(body.get('ua') or '')
    if page not in VISIT_PAGES or not VISIT_ID_RE.match(vid) or BOT_RE.search(ua):
        return
    today = date.today().isoformat()
    with _lock:
        conn = db()
        conn.execute('INSERT OR IGNORE INTO visit_daily (day) VALUES (?)', (today,))
        conn.execute(f'UPDATE visit_daily SET {page}_views = {page}_views + 1 WHERE day = ?', (today,))
        new = conn.execute('INSERT OR IGNORE INTO visit_seen (day, vid) VALUES (?, ?)', (today, vid)).rowcount
        if new:
            conn.execute('UPDATE visit_daily SET visitors = visitors + 1 WHERE day = ?', (today,))
            conn.execute('INSERT INTO visit_source (day, source, visitors) VALUES (?, ?, 1) '
                         'ON CONFLICT(day, source) DO UPDATE SET visitors = visitors + 1', (today, visit_source(body.get('ref'))))
            conn.execute('DELETE FROM visit_seen WHERE day < ?', ((date.today() - timedelta(days=1)).isoformat(),))
        conn.commit()


def visit_stats(days):
    days = max(7, min(int(days or 30), 400))
    today = date.today()
    start = (today - timedelta(days=days - 1)).isoformat()
    with _lock:
        conn = db()
        rows = {r['day']: dict(r) for r in conn.execute('SELECT * FROM visit_daily WHERE day >= ?', (start,))}
        src = conn.execute('SELECT source, SUM(visitors) n FROM visit_source WHERE day >= ? GROUP BY source ORDER BY n DESC',
                           (start,)).fetchall()
        total = conn.execute('SELECT COALESCE(SUM(visitors), 0) v, COALESCE(SUM(home_views + reserve_views), 0) pv, '
                             'MIN(day) since FROM visit_daily').fetchone()
        month = conn.execute('SELECT COALESCE(SUM(visitors), 0) FROM visit_daily WHERE day >= ?',
                             (today.replace(day=1).isoformat(),)).fetchone()[0]
    items = []
    for i in range(days):  # 방문이 없던 날도 0 으로 채움 (최근 날짜가 위)
        d = (today - timedelta(days=i)).isoformat()
        r = rows.get(d, {})
        items.append({'day': d, 'visitors': r.get('visitors', 0),
                      'home_views': r.get('home_views', 0), 'reserve_views': r.get('reserve_views', 0)})
    return {'items': items, 'sources': [dict(r) for r in src], 'month': month,
            'total': total['v'], 'total_views': total['pv'], 'since': total['since']}


def handle_public(method, path, query, body, ip):
    """손님용 API — 로그인 없이 사용. (상태코드, 응답) 반환"""
    q = lambda k, d='': query.get(k, [d])[0]
    if method == 'GET' and path == '/api/public/info':
        return 200, public_info()
    if method == 'GET' and path == '/api/public/calendar':
        return 200, month_calendar(q('month'))
    if method == 'GET' and path == '/api/public/search':
        rooms = search(q('checkin'), q('checkout'))
        with _lock:
            left = option_left(load_settings(), q('checkin'), q('checkout'))
        return 200, {'rooms': rooms, 'option_left': left}
    if method == 'POST' and path == '/api/public/quote':
        settings, prices = load_settings(), load_prices()
        quotes = [quote(it.get('room_id'), body.get('checkin'), body.get('checkout'), it.get('guests'),
                        it.get('options'), prices, settings['maxNights']) for it in body.get('items') or []]
        return 200, {'items': quotes, 'total': sum(x['total'] for x in quotes)}
    if method == 'POST' and path == '/api/public/reserve':
        if body.get('website'):  # 스팸 봇 차단용 숨김 칸
            return 400, {'error': '잘못된 요청'}
        if rate_limited('reserve:' + ip, 5, 3600):
            return 429, {'error': '예약 시도가 너무 많습니다. 전화로 문의해 주세요'}
        return 200, create_web_booking(body)
    if method == 'POST' and path == '/api/public/cancel':
        if rate_limited('lookup:' + ip, 20, 600):
            return 429, {'error': '잠시 후 다시 시도해 주세요'}
        return 200, customer_cancel(body.get('code'), body.get('phone'), body)
    if method == 'POST' and path == '/api/public/lookup':
        if rate_limited('lookup:' + ip, 20, 600):
            return 429, {'error': '잠시 후 다시 시도해 주세요'}
        return 200, lookup(body.get('code'), body.get('phone'))
    if method == 'POST' and path == '/api/public/visit':
        if not rate_limited('visit:' + ip, 300, 3600):
            record_visit(body)
        return 200, {'ok': True}
    return 404, {'error': '알 수 없는 요청'}


def list_filters(q):
    """목록 · 엑셀 공통 검색 조건 (주소의 ?from=&to=&status=... 를 해석)"""
    split = lambda k: [x for x in q(k).split(',') if x] or None
    return {'start': q('from') or None, 'end': q('to') or None, 'status': split('status'), 'q': q('q').strip(),
            'basis': q('basis') or 'stay', 'room': q('room') or None, 'refund_status': split('refund'),
            'cancel_by': q('cancel_by') or None}


def handle_admin(method, path, query, body):
    """관리자용 API (로그인 필요). (상태코드, 응답) 또는 ('file', 내용, 형식, 파일명)"""
    q = lambda k, d='': query.get(k, [d])[0]
    if method == 'GET':
        if path == '/api/bk/reservations':
            return 200, {'items': list_reservations(**list_filters(q))}
        if path == '/api/bk/customers':
            return 200, list_customers(q('view') or 'all', q('q').strip(), q('category'))
        if path == '/api/bk/customer':
            return 200, get_customer(q('phone'))
        if path == '/api/bk/sales/monthly':
            return 200, {'items': sales_monthly(int(q('year')), q('basis') or 'checkout', q('exclude') == '1',
                                                q('pay') or None)}
        if path == '/api/bk/sales/daily':
            return 200, {'items': sales_daily(int(q('year')), int(q('month')), q('basis') or 'checkout',
                                              q('exclude') == '1', q('pay') or None)}
        if path == '/api/bk/sales/calendar':
            return 200, money_calendar(q('month'))
        if path == '/api/bk/sales/cost':
            return 200, cost_report(int(q('year')))
        if path == '/api/bk/monthly':
            return 200, {'items': monthly_report(q('from'), q('to'), q('basis') or 'created')}
        if path == '/api/bk/options':
            return 200, options_calendar(q('month'))
        if path == '/api/bk/option-left':
            with _lock:
                return 200, option_left(load_settings(), q('checkin'), q('checkout'), q('exclude') or None)
        if path == '/api/bk/reservation':
            with _lock:
                r = get_reservation(int(q('id', '0')))
                group = group_rows(r['group_code']) if r['group_code'] else [r]
            return 200, {'item': r, 'group': group}
        if path == '/api/bk/visits':
            return 200, visit_stats(q('days', '30'))
        if path == '/api/bk/stats':
            return 200, stats(int(q('year') or date.today().year))
        if path == '/api/bk/settings':
            prices = load_prices()
            return 200, {'settings': load_settings(), 'peakPeriods': prices.get('peakPeriods', []),
                         'holidays': prices.get('holidays', []), 'sources': SOURCES, 'payMethods': PAY_METHODS,
                         'refundPolicy': refund_policy_text(load_settings()),
                         'tgTemplateDefaults': DEFAULT_SETTINGS['tgTemplates'],
                         'rooms': [{'id': r['id'], 'name': r['name'], 'std': r['std'], 'max': r['max'],
                                    'group': r['group']} for r in prices['rooms']],
                         'options': [{'i': o['i'], 'name': o['name'], 'detail': o.get('detail', ''),
                                      'price': o['amount']} for o in paid_options(prices)]}
        if path == '/api/bk/sms/log':
            return 200, {'items': sms_log()}
        if path == '/api/bk/sms/balance':
            return 200, sms_balance()
        if path == '/api/bk/telegram/chat':
            return 200, telegram_find_chat()
        if path == '/api/bk/myip':  # 알리고에 등록할 이 PC의 인터넷(공인) IP
            try:
                with urllib.request.urlopen('https://api.ipify.org', timeout=8) as resp:
                    return 200, {'ip': resp.read().decode().strip()}
            except Exception:
                raise BookingError('인터넷 IP를 확인할 수 없습니다')
        if path == '/api/bk/export':
            f = list_filters(q)
            data, ctype, ext = export(f.pop('start'), f.pop('end'), **f)
            return 'file', data, ctype, f'reservations_{q("from")}_{q("to")}.{ext}'
        if path == '/api/bk/backup':
            return 'file', backup_zip(), 'application/zip', f'dhc-backup-{date.today().isoformat()}.zip'
    if method == 'POST':
        if path == '/api/bk/quote':
            return 200, quote(body.get('room_id'), body.get('checkin'), body.get('checkout'),
                              body.get('guests'), body.get('options'))
        if path == '/api/bk/create':
            return 200, {'item': create_reservation(body)}
        if path == '/api/bk/update':
            return 200, {'item': update_reservation(int(body.get('id', 0)), body)}
        if path == '/api/bk/status':
            return 200, {'items': set_status(int(body.get('id', 0)), body.get('status'), body)}
        if path == '/api/bk/customer/save':
            return 200, save_customer(body)
        if path == '/api/bk/customers/action':
            return 200, {'count': customers_action(body.get('phones') or [], body.get('action'), body.get('category', ''))}
        if path == '/api/bk/customers/sms':
            return 200, {'sent': customers_sms(body.get('phones') or [], body.get('msg'),
                                                body.get('categoryAll') if body.get('all') else None)}
        if path == '/api/bk/refund-done':
            refund_done(body.get('ids') or [], body.get('done', True), body.get('sms', True))
            return 200, {'ok': True}
        if path == '/api/bk/block-range':
            return 200, block_range(body.get('room') or '', body.get('start'), body.get('end'), body.get('reason'))
        if path == '/api/bk/block-all':
            return 200, {'made': block_all(body.get('date'), body.get('reason'))}
        if path == '/api/bk/sms/bulk':
            if not body.get('kind') and not str(body.get('msg', '')).strip():
                raise BookingError('문자 내용을 입력하세요')
            return 200, {'sent': bulk_sms(body.get('ids') or [], body.get('kind'), str(body.get('msg', '')))}
        if path == '/api/bk/delete':
            delete_reservation(int(body.get('id', 0)))
            return 200, {'ok': True}
        if path == '/api/bk/settings':
            return 200, {'settings': save_settings(body)}
        if path == '/api/bk/calendar-rules':
            save_calendar_rules(body.get('peakPeriods') or [], body.get('holidays') or [])
            return 200, {'ok': True}
        if path == '/api/bk/sms/send':
            msg = str(body.get('msg', '')).strip()
            if not msg:
                raise BookingError('문자 내용을 입력하세요')
            ok, result = send_sms(body.get('phone'), msg, 'manual', body.get('reservation_id'))
            return (200 if ok else 400), ({'ok': True, 'result': result} if ok else {'error': result})
        if path == '/api/bk/telegram/test':
            ok, result = send_telegram('✅ 동해카라반펜션 예약 알림이 연결되었습니다.\n새 예약 · 입금 확인 · 손님 취소 · 미입금 자동취소를 여기로 알려드립니다.', 'test')
            return (200, {'ok': True}) if ok else (400, {'error': result})
        if path == '/api/bk/sms/preview':
            with _lock:
                r = get_reservation(int(body.get('id', 0)))
                rows = group_rows(r['group_code']) if r['group_code'] else [r]
            return 200, {'msg': render_template(load_settings()['templates'].get(body.get('kind'), ''), rows)}
    return 404, {'error': '알 수 없는 요청'}
